from datetime import date
from decimal import Decimal
import re
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from core.models import SiteSettings
from masters.models import Client, ClientGroup
from tasks.models import Task, TaskGroup, TaskMaster
from tasks.services import complete_task, create_task_from_master, submit_task, verify_task

from accounts.forms import AccountForm, InvoiceItemForm, InvoicePartyForm, JournalLineFormSet
from accounts.models import (
    SERIES_GST,
    SERIES_GST_B2C,
    SERIES_NGST,
    SERIES_PUR,
    VOUCHER_TABS,
    Account,
    AccountGroup,
    InvoiceNumberSeq,
    ReceiptAllocation,
    Voucher,
)
from accounts.export import (
    _build_pdf,
    _pdf_text,
    balance_sheet_pdf_bytes,
    balance_sheet_xlsx_bytes,
    billwise_outstanding_pdf_bytes,
    book_register_pdf_bytes,
    ledger_pdf_bytes,
    ledger_xlsx_bytes,
    profit_and_loss_pdf_bytes,
    profit_and_loss_xlsx_bytes,
    trial_balance_pdf_bytes,
    trial_balance_xlsx_bytes,
)
from accounts.services import (
    allocate_voucher_number,
    amount_in_indian_words,
    apply_receipt_allocations,
    account_totals,
    build_contra_lines,
    build_invoice_lines,
    build_payment_lines,
    build_receipt_lines,
    capital_ledgers,
    cash_equivalent_ledgers,
    line_account_ledgers,
    non_cash_ledgers,
    payment_account_ledgers,
    receipt_account_ledgers,
    task_master_ledgers,
    outstanding_bills_for_account,
    outstanding_party_bills_for_note,
    outstanding_purchase_bills_for_account,
    create_sales_invoice_for_task,
    delete_parent_and_account_groups,
    delete_voucher,
    post_voucher,
    sales_invoice_document,
    seed_chart_of_accounts,
    sync_client_account,
    TB_FORMAT_OPENING,
    TB_FORMAT_TB,
    TB_FORMAT_TRANSACTIONS,
    TB_REPORT_GROUP,
    TB_REPORT_LEDGER,
    all_ledger_sections,
    balance_sheet,
    billwise_outstanding_sections,
    book_register_rows,
    ledger_lines,
    sales_register_totals,
    profit_and_loss,
    trial_balance,
    trial_balance_drill,
)

User = get_user_model()


def _assert_pdf_structure(testcase, payload: bytes):
    testcase.assertTrue(payload.startswith(b"%PDF-"))
    testcase.assertIn(b"%%EOF", payload[-64:])
    obj_nums = {int(n) for n in re.findall(rb"^(\d+) 0 obj", payload, re.M)}
    testcase.assertTrue(obj_nums)
    for n in re.findall(rb"/Contents (\d+) 0 R", payload):
        num = int(n)
        testcase.assertIn(num, obj_nums)
        match = re.search(rf"{num} 0 obj(.*?)endobj".encode(), payload, re.S)
        testcase.assertIsNotNone(match)
        testcase.assertIn(b"stream", match.group(1))
    kids = re.search(rb"/Kids \[([^\]]*)\]", payload)
    testcase.assertIsNotNone(kids)
    for kn in re.findall(rb"(\d+) 0 R", kids.group(1)):
        testcase.assertIn(int(kn), obj_nums)


def _posting_ledgers():
    seed_chart_of_accounts()
    income_head = AccountGroup.objects.get(system_key="income")
    liabilities = AccountGroup.objects.get(system_key="liabilities")
    gst_group, _ = AccountGroup.objects.get_or_create(
        name="GST Output",
        parent=liabilities,
        defaults={"main_head": "liabilities"},
    )
    income, _ = Account.objects.get_or_create(
        system_key="professional_income",
        defaults={"name": "Professional income", "group": income_head},
    )
    for name, rate in (
        ("CGST Output 9%", Decimal("9.00")),
        ("SGST Output 9%", Decimal("9.00")),
        ("IGST Output 18%", Decimal("18.00")),
        ("CGST Output 2.5%", Decimal("2.50")),
        ("SGST Output 2.5%", Decimal("2.50")),
        ("IGST Output 5%", Decimal("5.00")),
    ):
        Account.objects.get_or_create(
            name=name,
            group=gst_group,
            defaults={"gst_rate": rate},
        )
    return income


class AccountsModuleTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="acc-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        self.user = User.objects.create_user(email="acc-user@example.com", password="pass12345")
        self.verifier = User.objects.create_user(email="acc-v@example.com", password="pass12345")
        self.doc = User.objects.create_user(email="acc-d@example.com", password="pass12345")
        grp = ClientGroup.objects.create(name="ACC GROUP")
        self.client_row = Client.objects.create(
            client_name="ACC CLIENT",
            client_type="Individual",
            branch="Trivandrum",
            client_group=grp,
            approval_status=Client.APPROVED,
            pan="ABCDE7777F",
        )
        tg = TaskGroup.objects.create(name="GST ACC", sort_order=1)
        self.master = TaskMaster.objects.create(
            task_group=tg,
            name="GSTR billing acc",
            is_recurring=False,
            requires_verifier=True,
            requires_document_checker=True,
        )

    def test_nested_group_inherits_main_head(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        child = AccountGroup(name="Fixed assets parent", parent=assets, main_head="income")
        child.full_clean()
        self.assertEqual(child.main_head, "assets")
        self.assertTrue(child.is_parent_group())
        self.assertEqual(AccountGroup.account_groups().count(), 0)

    def test_contra_tab_is_between_payment_and_journal(self):
        labels = [label for _key, label in VOUCHER_TABS]
        self.assertEqual(labels[labels.index("Payment") + 1], "Contra")
        self.assertEqual(labels[labels.index("Contra") + 1], "Journal")
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash = Account.objects.create(name="Contra cash", group=assets)
        bank = Account.objects.create(name="Contra bank", group=assets)
        voucher = post_voucher(
            voucher_type="contra",
            voucher_date=date(2026, 10, 16),
            lines=[
                {"account": bank, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": cash, "debit": Decimal("0.00"), "credit": Decimal("100.00")},
            ],
            actor=self.admin,
        )
        self.assertTrue(voucher.number.startswith("CV/"))

    def test_contra_debits_receiver_and_credits_payer(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        cash = Account.objects.create(name="Till cash", group=cash_group)
        bank = Account.objects.create(name="Current bank", group=cash_group)
        rent = Account.objects.create(
            name="Contra rent",
            group=AccountGroup.objects.get(system_key="expenses"),
        )
        form = InvoiceItemForm(voucher_type="contra")
        names = set(form.fields["account"].queryset.values_list("name", flat=True))
        self.assertIn("Till cash", names)
        self.assertIn("Current bank", names)
        self.assertNotIn("Contra rent", names)
        with self.assertRaises(ValidationError):
            build_contra_lines(
                receivers=[{"account": cash, "amount": Decimal("50.00")}],
                payers=[{"account": cash, "amount": Decimal("50.00")}],
            )
        with self.assertRaises(ValidationError):
            build_contra_lines(
                receivers=[{"account": bank, "amount": Decimal("75.00")}],
                payers=[{"account": cash, "amount": Decimal("50.00")}],
            )
        lines = build_contra_lines(
            receivers=[{"account": bank, "amount": Decimal("75.00")}],
            payers=[{"account": cash, "amount": Decimal("75.00")}],
        )
        self.assertEqual(lines[0]["account"], bank)
        self.assertEqual(lines[0]["debit"], Decimal("75.00"))
        self.assertEqual(lines[1]["account"], cash)
        self.assertEqual(lines[1]["credit"], Decimal("75.00"))
        self.client.force_login(self.admin)
        posted = self.client.post(
            reverse("accounts_voucher_create") + "?type=contra",
            {
                "voucher_type": "contra",
                "date": "2026-10-16",
                "narration": "Till to bank",
                "recv-TOTAL_FORMS": "1",
                "recv-INITIAL_FORMS": "0",
                "recv-MIN_NUM_FORMS": "0",
                "recv-MAX_NUM_FORMS": "1000",
                "recv-0-account": str(bank.pk),
                "recv-0-amount": "75.00",
                "pay-TOTAL_FORMS": "1",
                "pay-INITIAL_FORMS": "0",
                "pay-MIN_NUM_FORMS": "0",
                "pay-MAX_NUM_FORMS": "1000",
                "pay-0-account": str(cash.pk),
                "pay-0-amount": "75.00",
            },
            HTTP_HOST="localhost",
        )
        self.assertEqual(posted.status_code, 302)
        voucher = Voucher.objects.get(number__startswith="CV/")
        self.assertEqual(voucher.debit_total(), Decimal("75.00"))
        self.assertEqual(voucher.credit_total(), Decimal("75.00"))

    def test_voucher_date_uses_day_month_year(self):
        from accounts.forms import VoucherForm

        form = VoucherForm(data={"voucher_type": "journal", "date": "22-10-2026", "narration": ""})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["date"], date(2026, 10, 22))
        rendered = VoucherForm()
        rendered.fields["date"].initial = date(2026, 10, 22)
        html = str(rendered["date"])
        self.assertIn('value="22-10-2026"', html)
        self.assertIn("DD-MM-YYYY", html)

    def test_vouchers_can_be_deleted(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash = Account.objects.create(name="Delete cash", group=assets)
        bank = Account.objects.create(name="Delete bank", group=assets)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            narration="To be deleted",
            lines=[
                {"account": cash, "debit": Decimal("40.00"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("0.00"), "credit": Decimal("40.00")},
            ],
            actor=self.admin,
        )
        pk = voucher.pk
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("accounts_voucher_delete", args=[pk]),
            HTTP_HOST="localhost",
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Voucher.objects.filter(pk=pk).exists())
        self.assertFalse(Voucher.objects.filter(number=voucher.number).exists())

    def test_voucher_can_be_edited(self):
        from accounts.services import update_voucher

        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash = Account.objects.create(name="Edit cash", group=assets)
        bank = Account.objects.create(name="Edit bank", group=assets)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            narration="Original",
            lines=[
                {"account": cash, "debit": Decimal("40.00"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("0.00"), "credit": Decimal("40.00")},
            ],
            actor=self.admin,
        )
        number = voucher.number
        update_voucher(
            voucher,
            voucher_date=date(2026, 10, 18),
            narration="Updated",
            lines=[
                {"account": cash, "debit": Decimal("55.00"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("0.00"), "credit": Decimal("55.00")},
            ],
        )
        voucher.refresh_from_db()
        self.assertEqual(voucher.number, number)
        self.assertEqual(voucher.date, date(2026, 10, 18))
        self.assertEqual(voucher.narration, "Updated")
        self.assertEqual(voucher.debit_total(), Decimal("55.00"))
        lines = list(voucher.lines.order_by("pk"))
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("accounts_voucher_edit", args=[voucher.pk]),
            {
                "voucher_type": "journal",
                "date": "2026-10-19",
                "narration": "Posted edit",
                "lines-TOTAL_FORMS": "3",
                "lines-INITIAL_FORMS": "2",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-id": str(lines[0].pk),
                "lines-0-account": str(cash.pk),
                "lines-0-debit": "70.00",
                "lines-0-credit": "0.00",
                "lines-1-id": str(lines[1].pk),
                "lines-1-account": str(bank.pk),
                "lines-1-debit": "0.00",
                "lines-1-credit": "70.00",
                "lines-2-account": "",
                "lines-2-debit": "0.00",
                "lines-2-credit": "0.00",
            },
            HTTP_HOST="localhost",
        )
        self.assertEqual(response.status_code, 302)
        voucher.refresh_from_db()
        self.assertEqual(voucher.number, number)
        self.assertEqual(voucher.date, date(2026, 10, 19))
        self.assertEqual(voucher.narration, "Posted edit")
        self.assertEqual(voucher.debit_total(), Decimal("70.00"))

    def test_voucher_list_pdf_export(self):
        from accounts.export import voucher_register_pdf_bytes

        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash = Account.objects.create(name="List cash", group=assets)
        bank = Account.objects.create(name="List bank", group=assets)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            narration="List export narration",
            lines=[
                {"account": cash, "debit": Decimal("125.50"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("0.00"), "credit": Decimal("125.50")},
            ],
            actor=self.admin,
        )
        pdf = voucher_register_pdf_bytes(
            [voucher],
            title="Journal voucher",
            company_name="Test Firm",
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Date", pdf)
        self.assertIn(b"Voucher No", pdf)
        self.assertIn(b"Narration", pdf)
        self.assertIn(b"Amount", pdf)
        self.assertIn(voucher.number.encode(), pdf)
        self.assertIn(b"List export narration", pdf)
        self.assertIn(b"125.50", pdf)
        self.client.force_login(self.admin)
        response = self.client.get(
            reverse("accounts_voucher_list"),
            {"type": "journal", "export": "pdf"},
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("vouchers-journal.pdf", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(voucher.number.encode(), response.content)

    def test_sales_voucher_list_is_latest_month_with_client_and_totals(self):
        from accounts.views import _latest_voucher_month, _voucher_register_qs

        seed_chart_of_accounts()
        income = AccountGroup.objects.get(system_key="income")
        assets = AccountGroup.objects.get(system_key="assets")
        fees = Account.objects.create(name="List fees", group=income)
        party = Account.objects.create(name="Pearl Generators List", group=assets)
        older = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 9, 10),
            narration="September sale",
            lines=[
                {"account": party, "debit": Decimal("1180.00"), "credit": Decimal("0.00")},
                {"account": fees, "debit": Decimal("0.00"), "credit": Decimal("1180.00")},
            ],
            actor=self.admin,
        )
        latest = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 7),
            narration="October sale",
            lines=[
                {"account": party, "debit": Decimal("2360.00"), "credit": Decimal("0.00")},
                {"account": fees, "debit": Decimal("0.00"), "credit": Decimal("2360.00")},
            ],
            actor=self.admin,
        )
        date_from, date_to = _latest_voucher_month("sales")
        self.assertEqual(date_from, date(2026, 10, 1))
        self.assertEqual(date_to, date(2026, 10, 31))
        rows = list(_voucher_register_qs("sales", date_from=date_from, date_to=date_to))
        self.assertEqual([row.pk for row in rows], [latest.pk])
        self.assertEqual(rows[0].client_name, "Pearl Generators List")
        self.assertEqual(rows[0].debit_amount, Decimal("2360.00"))
        self.assertEqual(rows[0].credit_amount, Decimal("2360.00"))
        self.assertNotIn(older.pk, [row.pk for row in rows])
        from accounts.export import voucher_register_pdf_bytes

        pdf = voucher_register_pdf_bytes(rows, title="Sales invoice — October 2026", voucher_type="sales")
        self.assertIn(b"Client name", pdf)
        self.assertIn(b"Pearl Generators List", pdf)
        self.assertIn(b"Debit", pdf)
        self.assertIn(b"Credit", pdf)

    def test_book_register_lists_and_exports(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash = Account.objects.create(name="Register cash", group=assets)
        bank = Account.objects.create(name="Register bank", group=assets)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            narration="Register narration",
            lines=[
                {"account": cash, "debit": Decimal("80.00"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("0.00"), "credit": Decimal("80.00")},
            ],
            actor=self.admin,
        )
        rows = book_register_rows(
            voucher_type="journal",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["number"], voucher.number)
        self.assertEqual(rows[0]["particulars"], "Register cash")
        self.assertEqual(rows[0]["amount"], Decimal("80.00"))
        pdf = book_register_pdf_bytes(
            rows,
            title="Journal Register",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Particulars", pdf)
        self.assertIn(voucher.number.encode(), pdf)
        self.client.force_login(self.admin)
        pdf_res = self.client.get(
            reverse("accounts_register"),
            {"type": "journal", "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "pdf"},
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")
        xlsx_res = self.client.get(
            reverse("accounts_register"),
            {"type": "journal", "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "xlsx"},
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_res["Content-Type"])

    def test_sales_register_gst_and_ngst_formats(self):
        income = _posting_ledgers()
        income.gst_rate = Decimal("18.00")
        income.save(update_fields=["gst_rate"])
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        b2b_client = Client.objects.create(
            client_name="B2B Buyer",
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.client_row.client_group,
            approval_status=Client.APPROVED,
            pan="ABCDE1111F",
            gstin="32ABCDE1111F1Z5",
        )
        b2c_client = Client.objects.create(
            client_name="B2C Buyer",
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.client_row.client_group,
            approval_status=Client.APPROVED,
            pan="ABCDE2222F",
        )
        ngst_client = Client.objects.create(
            client_name="NGST Buyer",
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.client_row.client_group,
            approval_status=Client.APPROVED,
            pan="ABCDE3333F",
        )
        b2b = sync_client_account(b2b_client)
        b2c = sync_client_account(b2c_client)
        ngst_party = sync_client_account(ngst_client)
        expenses_head = AccountGroup.objects.get(system_key="expenses")
        recovery = Account.objects.create(
            name="ROC recovery",
            group=expenses_head,
            is_pure_agent=True,
        )
        post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 15),
            narration="B2B GST",
            lines=build_invoice_lines(
                voucher_type="sales",
                party=b2b,
                items=[
                    {"account": income, "amount": Decimal("1000.00")},
                    {"account": recovery, "amount": Decimal("200.00")},
                ],
                with_gst=True,
            ),
            actor=self.admin,
            with_gst=True,
        )
        post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 16),
            narration="B2C GST",
            lines=build_invoice_lines(
                voucher_type="sales",
                party=b2c,
                items=[{"account": income, "amount": Decimal("500.00")}],
                with_gst=True,
            ),
            actor=self.admin,
            with_gst=True,
        )
        post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 16),
            narration="Non GST bill",
            lines=build_invoice_lines(
                voucher_type="sales",
                party=ngst_party,
                items=[{"account": income, "amount": Decimal("800.00")}],
            ),
            actor=self.admin,
            with_gst=False,
        )
        gst_rows = book_register_rows(
            voucher_type="sales",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 31),
            series=SERIES_GST,
        )
        self.assertEqual(len(gst_rows), 2)
        b2b_row = next(row for row in gst_rows if row["gstin"] == "32ABCDE1111F1Z5")
        b2c_row = next(row for row in gst_rows if not row["gstin"])
        self.assertEqual(b2b_row["customer"], "B2B BUYER")
        self.assertEqual(b2b_row["b2b_sales"], Decimal("1000.00"))
        self.assertEqual(b2b_row["b2c_sales"], Decimal("0.00"))
        self.assertEqual(b2b_row["pure_agent"], Decimal("200.00"))
        self.assertEqual(b2b_row["cgst"], Decimal("90.00"))
        self.assertEqual(b2b_row["sgst"], Decimal("90.00"))
        self.assertEqual(b2b_row["igst"], Decimal("0.00"))
        self.assertEqual(b2b_row["total"], Decimal("1380.00"))
        self.assertEqual(b2c_row["customer"], "B2C BUYER")
        self.assertEqual(b2c_row["b2c_sales"], Decimal("500.00"))
        self.assertEqual(b2c_row["b2b_sales"], Decimal("0.00"))
        self.assertEqual(b2c_row["total"], Decimal("590.00"))
        gst_totals = sales_register_totals(gst_rows, gst=True)
        self.assertEqual(gst_totals["b2b_sales"], Decimal("1000.00"))
        self.assertEqual(gst_totals["b2c_sales"], Decimal("500.00"))
        self.assertEqual(gst_totals["total"], Decimal("1970.00"))
        ngst_rows = book_register_rows(
            voucher_type="sales",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 31),
            series=SERIES_NGST,
        )
        self.assertEqual(len(ngst_rows), 1)
        self.assertEqual(ngst_rows[0]["customer"], "NGST BUYER")
        self.assertEqual(ngst_rows[0]["sales"], Decimal("800.00"))
        self.assertEqual(ngst_rows[0]["total"], Decimal("800.00"))
        gst_pdf = book_register_pdf_bytes(
            gst_rows,
            title="Sales register — GST Invoice",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 31),
            invoice_kind="gst",
        )
        self.assertTrue(gst_pdf.startswith(b"%PDF"))
        self.assertIn(b"B2B Sales", gst_pdf)
        self.assertIn(b"GSTIN NO", gst_pdf)
        self.assertNotIn(b"Particulars", gst_pdf)
        ngst_pdf = book_register_pdf_bytes(
            ngst_rows,
            title="Sales register — Non GST Invoice",
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 31),
            invoice_kind="ngst",
        )
        self.assertTrue(ngst_pdf.startswith(b"%PDF"))
        self.assertIn(b"Customer Name", ngst_pdf)
        self.client.force_login(self.admin)
        gst_xlsx = self.client.get(
            reverse("accounts_register"),
            {
                "type": "sales",
                "invoice": "gst",
                "date_from": "2026-04-01",
                "date_to": "2026-10-31",
                "export": "xlsx",
            },
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(gst_xlsx.status_code, 200)

    def test_billwise_outstanding_groups_debtor_bills(self):
        income = _posting_ledgers()
        buyer = sync_client_account(self.client_row)
        debtors = AccountGroup.objects.create(
            name="Sundry Debtors",
            parent=AccountGroup.objects.get(system_key="assets"),
            main_head="assets",
        )
        buyer.group = debtors
        buyer.save(update_fields=["group"])
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 1),
            lines=build_invoice_lines(
                voucher_type="sales",
                party=buyer,
                items=[{"account": income, "amount": Decimal("1000.00")}],
            ),
            actor=self.admin,
        )
        post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 10),
            lines=build_receipt_lines(
                mode=Account.objects.create(
                    name="Outstanding cash",
                    group=AccountGroup.objects.get(system_key="assets"),
                ),
                items=[{"account": buyer, "amount": Decimal("400.00")}],
            ),
            actor=self.admin,
            receipt_items=[
                {
                    "account": buyer,
                    "amount": Decimal("400.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": invoice.pk, "amount": "400.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        data = billwise_outstanding_sections(as_on=date(2026, 10, 16))
        self.assertEqual(len(data["sections"]), 1)
        section = data["sections"][0]
        self.assertEqual(section["name"], buyer.name)
        self.assertEqual(len(section["rows"]), 1)
        self.assertEqual(section["rows"][0]["sl_no"], 1)
        self.assertEqual(section["rows"][0]["number"], invoice.number)
        self.assertEqual(section["rows"][0]["outstanding"], Decimal("600.00"))
        self.assertEqual(section["rows"][0]["due_for"], "15 Days")
        self.assertEqual(data["grand_total"], Decimal("600.00"))
        pdf = billwise_outstanding_pdf_bytes(data, title="Bill wise outstanding", company_name="Test Firm")
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Total Outstanding", pdf)
        self.assertIn(b"Due For", pdf)
        self.client.force_login(self.admin)
        xlsx = self.client.get(
            reverse("accounts_billwise_outstanding"),
            {"date_to": "2026-10-16", "export": "xlsx"},
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(xlsx.status_code, 200)

    def test_deleting_invoice_unwinds_receipt_knock_off(self):
        income = _posting_ledgers()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash delete group",
            parent=assets,
            main_head="assets",
        )
        bank = Account.objects.create(name="Delete receipt bank", group=cash_group)
        buyer = sync_client_account(self.client_row)
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 1),
            lines=build_invoice_lines(
                voucher_type="sales",
                party=buyer,
                items=[{"account": income, "amount": Decimal("1000.00")}],
            ),
            actor=self.admin,
        )
        receipt = post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 15),
            lines=build_receipt_lines(
                mode=bank,
                items=[{"account": buyer, "amount": Decimal("400.00")}],
            ),
            actor=self.admin,
            receipt_items=[
                {
                    "account": buyer,
                    "amount": Decimal("400.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": invoice.pk, "amount": "400.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertTrue(ReceiptAllocation.objects.filter(receipt=receipt, against=invoice).exists())
        delete_voucher(invoice)
        self.assertFalse(Voucher.objects.filter(pk=invoice.pk).exists())
        self.assertFalse(ReceiptAllocation.objects.filter(against_id=invoice.pk).exists())
        self.assertTrue(Voucher.objects.filter(pk=receipt.pk).exists())

    def test_journal_starts_with_two_lines_and_rejects_unequal_totals(self):
        seed_chart_of_accounts()
        cash = Account.objects.create(
            name="JV cash",
            group=AccountGroup.objects.get(system_key="assets"),
        )
        expense = Account.objects.create(
            name="JV expense",
            group=AccountGroup.objects.get(system_key="expenses"),
        )
        formset = JournalLineFormSet()
        self.assertEqual(formset.extra, 2)
        self.assertEqual(len(formset.forms), 2)
        with self.assertRaises(ValidationError):
            post_voucher(
                voucher_type="journal",
                voucher_date=date(2026, 10, 16),
                lines=[
                    {"account": expense, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                    {"account": cash, "debit": Decimal("0.00"), "credit": Decimal("40.00")},
                ],
                actor=self.admin,
            )
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            lines=[
                {"account": expense, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": cash, "debit": Decimal("0.00"), "credit": Decimal("100.00")},
            ],
            actor=self.admin,
        )
        self.assertTrue(voucher.number.startswith("JV/"))
        self.assertEqual(voucher.debit_total(), voucher.credit_total())
        balanced = [
            {"account": expense, "debit": Decimal("50.00"), "credit": Decimal("0.00")},
            {"account": cash, "debit": Decimal("0.00"), "credit": Decimal("50.00")},
        ]
        credit_note = post_voucher(
            voucher_type="credit_note",
            voucher_date=date(2026, 10, 16),
            lines=balanced,
            actor=self.admin,
        )
        self.assertTrue(credit_note.number.startswith("CN/"))
        debit_note = post_voucher(
            voucher_type="debit_note",
            voucher_date=date(2026, 10, 16),
            lines=balanced,
            actor=self.admin,
        )
        self.assertTrue(debit_note.number.startswith("DN/"))

    def test_credit_and_debit_notes_knock_off_party_bills(self):
        income = _posting_ledgers()
        buyer = sync_client_account(self.client_row)
        debtors = AccountGroup.objects.create(
            name="Sundry Debtors",
            parent=AccountGroup.objects.get(system_key="assets"),
            main_head="assets",
        )
        buyer.group = debtors
        buyer.save(update_fields=["group"])
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 1),
            lines=build_invoice_lines(
                voucher_type="sales",
                party=buyer,
                items=[{"account": income, "amount": Decimal("1000.00")}],
            ),
            actor=self.admin,
        )
        self.assertEqual(outstanding_party_bills_for_note(buyer)[0]["outstanding"], Decimal("1000.00"))
        self.assertEqual(outstanding_party_bills_for_note(income), [])
        credit_note = post_voucher(
            voucher_type="credit_note",
            voucher_date=date(2026, 10, 16),
            lines=[
                {"account": income, "debit": Decimal("400.00"), "credit": Decimal("0.00")},
                {"account": buyer, "debit": Decimal("0.00"), "credit": Decimal("400.00")},
            ],
            actor=self.admin,
            receipt_items=[
                {
                    "account": buyer,
                    "amount": Decimal("400.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": invoice.pk, "amount": "400.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_bills_for_account(buyer)[0]["outstanding"], Decimal("600.00"))
        knock = ReceiptAllocation.objects.get(receipt=credit_note, is_advance=False)
        self.assertEqual(knock.against_id, invoice.pk)
        expenses = AccountGroup.objects.get(system_key="expenses")
        expense = Account.objects.create(name="Stationery", group=expenses)
        cred_group = AccountGroup.objects.create(
            name="Sundry Creditors",
            parent=AccountGroup.objects.get(system_key="liabilities"),
            main_head="liabilities",
        )
        supplier = Account.objects.create(name="Paper mart", group=cred_group)
        purchase = post_voucher(
            voucher_type="purchase",
            voucher_date=date(2026, 10, 2),
            lines=build_invoice_lines(
                voucher_type="purchase",
                party=supplier,
                items=[{"account": expense, "amount": Decimal("500.00")}],
            ),
            actor=self.admin,
        )
        self.assertEqual(outstanding_purchase_bills_for_account(supplier)[0]["outstanding"], Decimal("500.00"))
        self.assertEqual(outstanding_party_bills_for_note(supplier)[0]["id"], purchase.pk)
        debit_note = post_voucher(
            voucher_type="debit_note",
            voucher_date=date(2026, 10, 16),
            lines=[
                {"account": supplier, "debit": Decimal("200.00"), "credit": Decimal("0.00")},
                {"account": expense, "debit": Decimal("0.00"), "credit": Decimal("200.00")},
            ],
            actor=self.admin,
            receipt_items=[
                {
                    "account": supplier,
                    "amount": Decimal("200.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": purchase.pk, "amount": "200.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_purchase_bills_for_account(supplier)[0]["outstanding"], Decimal("300.00"))
        self.assertEqual(ReceiptAllocation.objects.get(receipt=debit_note).against_id, purchase.pk)

    def test_journal_knocks_off_debtor_and_creditor_bills(self):
        income = _posting_ledgers()
        buyer = sync_client_account(self.client_row)
        debtors = AccountGroup.objects.create(
            name="Sundry Debtors",
            parent=AccountGroup.objects.get(system_key="assets"),
            main_head="assets",
        )
        buyer.group = debtors
        buyer.save(update_fields=["group"])
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 1),
            lines=build_invoice_lines(
                voucher_type="sales",
                party=buyer,
                items=[{"account": income, "amount": Decimal("1000.00")}],
            ),
            actor=self.admin,
        )
        jv_debtor = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 16),
            lines=[
                {"account": income, "debit": Decimal("250.00"), "credit": Decimal("0.00")},
                {"account": buyer, "debit": Decimal("0.00"), "credit": Decimal("250.00")},
            ],
            actor=self.admin,
            receipt_items=[
                {
                    "account": buyer,
                    "amount": Decimal("250.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": invoice.pk, "amount": "250.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_bills_for_account(buyer)[0]["outstanding"], Decimal("750.00"))
        self.assertEqual(ReceiptAllocation.objects.get(receipt=jv_debtor).against_id, invoice.pk)
        expenses = AccountGroup.objects.get(system_key="expenses")
        expense = Account.objects.create(name="Journal stationery", group=expenses)
        cred_group = AccountGroup.objects.create(
            name="Sundry Creditors",
            parent=AccountGroup.objects.get(system_key="liabilities"),
            main_head="liabilities",
        )
        supplier = Account.objects.create(name="Journal paper mart", group=cred_group)
        purchase = post_voucher(
            voucher_type="purchase",
            voucher_date=date(2026, 10, 2),
            lines=build_invoice_lines(
                voucher_type="purchase",
                party=supplier,
                items=[{"account": expense, "amount": Decimal("400.00")}],
            ),
            actor=self.admin,
        )
        jv_creditor = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 17),
            lines=[
                {"account": supplier, "debit": Decimal("150.00"), "credit": Decimal("0.00")},
                {"account": expense, "debit": Decimal("0.00"), "credit": Decimal("150.00")},
            ],
            actor=self.admin,
            receipt_items=[
                {
                    "account": supplier,
                    "amount": Decimal("150.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": purchase.pk, "amount": "150.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_purchase_bills_for_account(supplier)[0]["outstanding"], Decimal("250.00"))
        self.assertEqual(ReceiptAllocation.objects.get(receipt=jv_creditor).against_id, purchase.pk)

    def test_delete_parent_and_account_groups_clears_nested_rows(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        parent = AccountGroup.objects.create(name="Current assets", parent=assets, main_head="assets")
        debtors = AccountGroup.objects.create(name="Sundry Debtors", parent=parent, main_head="assets")
        ledger = Account.objects.create(name="Temp debtor", group=debtors)
        deleted = delete_parent_and_account_groups()
        self.assertGreaterEqual(deleted, 2)
        self.assertEqual(AccountGroup.parent_groups().count(), 0)
        self.assertEqual(AccountGroup.account_groups().count(), 0)
        ledger.refresh_from_db()
        self.assertEqual(ledger.group.system_key, "assets")

    def test_client_syncs_ledger(self):
        acc = sync_client_account(self.client_row)
        self.assertIsNotNone(acc)
        self.assertEqual(acc.name, self.client_row.client_name)
        self.assertEqual(acc.pan, self.client_row.pan)
        self.assertEqual(acc.group.name, "Sundry Debtors")
        self.assertTrue(acc.group.is_account_group())
        self.assertEqual(acc.group.parent.name, "Current Assets")
        self.assertEqual(acc.client_id, self.client_row.pk)

    def test_gst_and_ngst_number_series_are_independent(self):
        d = date(2026, 10, 15)
        n1 = allocate_voucher_number(SERIES_GST, d)
        n2 = allocate_voucher_number(SERIES_NGST, d)
        n3 = allocate_voucher_number(SERIES_GST, d)
        self.assertEqual(n1, "GST/2026/10/00001")
        self.assertEqual(n2, "NGST/2026/10/00001")
        self.assertEqual(n3, "GST/2026/10/00002")
        self.assertEqual(allocate_voucher_number(SERIES_GST_B2C, d), "GSTB2C/2026/10/00001")
        self.assertEqual(allocate_voucher_number(SERIES_PUR, d), "PUR/2026/10/00001")

    def test_sales_gst_b2c_uses_separate_series(self):
        income = _posting_ledgers()
        income.gst_rate = Decimal("18.00")
        income.save(update_fields=["gst_rate"])
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        buyer = sync_client_account(self.client_row)
        form = InvoicePartyForm(voucher_type="sales")
        self.assertEqual(
            [key for key, _label in form.fields["tax_kind"].choices],
            ["gst", "gst_b2c", "ngst"],
        )
        lines = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[{"account": income, "amount": Decimal("1000.00")}],
            with_gst=True,
        )
        voucher = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 15),
            lines=lines,
            actor=self.admin,
            tax_kind="gst_b2c",
        )
        self.assertEqual(voucher.series, SERIES_GST_B2C)
        self.assertTrue(voucher.number.startswith("GSTB2C/"))
        b2b = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 15),
            lines=lines,
            actor=self.admin,
            tax_kind="gst",
        )
        self.assertEqual(b2b.series, SERIES_GST)
        self.assertTrue(b2b.number.startswith("GST/"))
        b2c_rows = book_register_rows(
            voucher_type="sales",
            date_from=date(2026, 10, 1),
            date_to=date(2026, 10, 31),
            series=SERIES_GST_B2C,
        )
        self.assertEqual([row["number"] for row in b2c_rows], [voucher.number])

    def test_sales_and_purchase_party_amounts(self):
        income = _posting_ledgers()
        buyer = sync_client_account(self.client_row)
        sales_lines = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[{"account": income, "amount": Decimal("1000.00")}],
        )
        self.assertEqual(sales_lines[0]["account"], buyer)
        self.assertEqual(sales_lines[0]["debit"], Decimal("1000.00"))
        self.assertEqual(sales_lines[1]["account"], income)
        self.assertEqual(sales_lines[1]["credit"], Decimal("1000.00"))
        income.gst_rate = Decimal("18.00")
        income.hsn_code = "998314"
        income.save(update_fields=["gst_rate", "hsn_code"])
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        buyer.gstin = "32ABCDE7777F1Z5"
        buyer.save(update_fields=["gstin"])
        gst_lines = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[{"account": income, "amount": Decimal("1000.00")}],
            with_gst=True,
        )
        by_name = {r["account"].name: r["credit"] for r in gst_lines}
        self.assertEqual(gst_lines[0]["debit"], Decimal("1180.00"))
        self.assertEqual(gst_lines[1]["credit"], Decimal("1000.00"))
        self.assertEqual(gst_lines[1]["hsn_code"], "998314")
        self.assertEqual(by_name["CGST Output 9%"], Decimal("90.00"))
        self.assertEqual(by_name["SGST Output 9%"], Decimal("90.00"))
        self.assertNotIn("IGST Output 18%", by_name)
        buyer.gstin = "33ABCDE7777F1Z5"
        buyer.save(update_fields=["gstin"])
        igst_lines = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[{"account": income, "amount": Decimal("1000.00")}],
            with_gst=True,
        )
        igst_by_name = {r["account"].name: r["credit"] for r in igst_lines}
        self.assertEqual(igst_lines[0]["debit"], Decimal("1180.00"))
        self.assertEqual(igst_by_name["IGST Output 18%"], Decimal("180.00"))
        self.assertNotIn("CGST Output 9%", igst_by_name)
        buyer.gstin = "32ABCDE7777F1Z5"
        buyer.save(update_fields=["gstin"])
        expenses_head = AccountGroup.objects.get(system_key="expenses")
        recovery = Account.objects.create(
            name="Reimbursement ROC",
            group=expenses_head,
            is_pure_agent=True,
            gst_rate=Decimal("18.00"),
        )
        mixed = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[
                {"account": income, "amount": Decimal("1000.00")},
                {"account": recovery, "amount": Decimal("500.00")},
            ],
            with_gst=True,
        )
        self.assertEqual(mixed[0]["debit"], Decimal("1680.00"))
        self.assertTrue(any(r["account"] == recovery and r["credit"] == Decimal("500.00") and r["is_pure_agent"] for r in mixed))
        self.assertEqual(
            sum(
                (r["credit"] for r in mixed if r["account"].name in ("CGST Output 9%", "SGST Output 9%")),
                Decimal("0"),
            ),
            Decimal("180.00"),
        )
        expenses = AccountGroup.objects.get(system_key="expenses")
        rent = Account.objects.create(name="Office rent", group=expenses)
        liabilities = AccountGroup.objects.get(system_key="liabilities")
        supplier = Account.objects.create(name="ABC Traders", group=liabilities)
        purchase_lines = build_invoice_lines(
            voucher_type="purchase",
            party=supplier,
            items=[{"account": rent, "amount": Decimal("250.00")}],
        )
        self.assertEqual(purchase_lines[0]["account"], supplier)
        self.assertEqual(purchase_lines[0]["credit"], Decimal("250.00"))
        self.assertEqual(purchase_lines[1]["account"], rent)
        self.assertEqual(purchase_lines[1]["debit"], Decimal("250.00"))
        voucher = post_voucher(
            voucher_type="purchase",
            voucher_date=date(2026, 10, 15),
            lines=purchase_lines,
            actor=self.admin,
        )
        self.assertTrue(voucher.number.startswith("PUR/"))
        self.assertEqual(voucher.debit_total(), voucher.credit_total())

    def test_sales_invoice_from_completed_task(self):
        _posting_ledgers()
        sync_client_account(self.client_row)
        task = create_task_from_master(
            master=self.master,
            client=self.client_row,
            assignee_users=[self.user],
            verifier=self.verifier,
            document_checker=self.doc,
            created_by=self.admin,
            period_key="2026-09",
            due_date=date(2026, 9, 20),
            is_billable=True,
            fees_amount=Decimal("1000.00"),
            auto_created=True,
        )
        submit_task(task, self.user)
        verify_task(task, self.verifier)
        complete_task(task, self.doc)
        task.refresh_from_db()
        voucher = create_sales_invoice_for_task(
            task=task, actor=self.admin, voucher_date=date(2026, 10, 1), with_gst=False
        )
        task.refresh_from_db()
        self.assertTrue(task.bill_raised)
        self.assertTrue(voucher.number.startswith("NGST/"))
        self.assertEqual(voucher.debit_total(), voucher.credit_total())
        income_line = next(ln for ln in voucher.lines.all() if ln.credit > 0)
        self.assertEqual(income_line.account.name, task.display_title)
        self.assertEqual(income_line.account.group.name, self.master.name)
        self.assertEqual(income_line.account.group.parent.name, self.master.task_group.name)
        with self.assertRaises(ValidationError):
            create_sales_invoice_for_task(task=task, actor=self.admin)

    def test_task_bill_reuses_existing_gst_filing_ledger_rate(self):
        _posting_ledgers()
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"company_name": "Test Firm", "gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        self.client_row.gstin = "32ABCDE7777F1Z5"
        self.client_row.save(update_fields=["gstin"])
        income_head = AccountGroup.objects.get(system_key="income")
        existing = Account.objects.create(
            name="GST Filing",
            group=income_head,
            gst_rate=Decimal("18.00"),
        )
        task = create_task_from_master(
            master=self.master,
            client=self.client_row,
            assignee_users=[self.user],
            verifier=self.verifier,
            document_checker=self.doc,
            created_by=self.admin,
            period_key="2026-06",
            due_date=date(2026, 6, 20),
            is_billable=True,
            fees_amount=Decimal("1000.00"),
            auto_created=True,
        )
        task.title = "GST Filing"
        task.save(update_fields=["title"])
        submit_task(task, self.user)
        verify_task(task, self.verifier)
        complete_task(task, self.doc)
        voucher = create_sales_invoice_for_task(
            task=task, actor=self.admin, voucher_date=date(2026, 10, 3), with_gst=True
        )
        income_line = next(ln for ln in voucher.lines.all() if ln.credit == Decimal("1000.00"))
        self.assertEqual(income_line.account_id, existing.pk)
        by_name = {ln.account.name: ln.credit for ln in voucher.lines.all() if ln.credit > 0}
        self.assertEqual(by_name["CGST Output 9%"], Decimal("90.00"))
        self.assertEqual(by_name["SGST Output 9%"], Decimal("90.00"))

    def test_unbilled_gst_series_follows_client_and_company_gstin(self):
        income = _posting_ledgers()
        income.gst_rate = Decimal("18.00")
        income.save(update_fields=["gst_rate"])
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"company_name": "Test Firm", "gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        self.client_row.gstin = "32ABCDE7777F1Z5"
        self.client_row.save(update_fields=["gstin"])
        intra_task = create_task_from_master(
            master=self.master,
            client=self.client_row,
            assignee_users=[self.user],
            verifier=self.verifier,
            document_checker=self.doc,
            created_by=self.admin,
            period_key="2026-07",
            due_date=date(2026, 7, 20),
            is_billable=True,
            fees_amount=Decimal("1000.00"),
            auto_created=True,
        )
        submit_task(intra_task, self.user)
        verify_task(intra_task, self.verifier)
        complete_task(intra_task, self.doc)
        voucher = create_sales_invoice_for_task(
            task=intra_task, actor=self.admin, voucher_date=date(2026, 10, 1), with_gst=True
        )
        self.assertTrue(voucher.number.startswith("GST/"))
        by_name = {ln.account.name: ln.credit for ln in voucher.lines.all() if ln.credit > 0}
        self.assertEqual(by_name[intra_task.display_title], Decimal("1000.00"))
        self.assertEqual(by_name["CGST Output 9%"], Decimal("90.00"))
        self.assertEqual(by_name["SGST Output 9%"], Decimal("90.00"))
        self.assertNotIn("IGST Output 18%", by_name)
        party_debit = next(ln.debit for ln in voucher.lines.all() if ln.debit > 0)
        self.assertEqual(party_debit, Decimal("1180.00"))

        self.client_row.gstin = "33ABCDE7777F1Z5"
        self.client_row.save(update_fields=["gstin"])
        inter_task = create_task_from_master(
            master=self.master,
            client=self.client_row,
            assignee_users=[self.user],
            verifier=self.verifier,
            document_checker=self.doc,
            created_by=self.admin,
            period_key="2026-08",
            due_date=date(2026, 8, 20),
            is_billable=True,
            fees_amount=Decimal("1000.00"),
            auto_created=True,
        )
        submit_task(inter_task, self.user)
        verify_task(inter_task, self.verifier)
        complete_task(inter_task, self.doc)
        igst_voucher = create_sales_invoice_for_task(
            task=inter_task, actor=self.admin, voucher_date=date(2026, 10, 2), with_gst=True
        )
        igst_by_name = {ln.account.name: ln.credit for ln in igst_voucher.lines.all() if ln.credit > 0}
        self.assertEqual(igst_by_name["IGST Output 18%"], Decimal("180.00"))
        self.assertNotIn("CGST Output 9%", igst_by_name)

    def test_invoice_first_marks_open_task_bill_raised(self):
        income = _posting_ledgers()
        party = sync_client_account(self.client_row)
        task = create_task_from_master(
            master=self.master,
            client=self.client_row,
            assignee_users=[self.user],
            verifier=self.verifier,
            document_checker=self.doc,
            created_by=self.admin,
            period_key="2026-10",
            due_date=date(2026, 10, 20),
            is_billable=True,
            fees_amount=Decimal("500.00"),
            auto_created=True,
        )
        voucher = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 2),
            narration="Advance bill",
            lines=[
                {"account": party, "debit": Decimal("500.00"), "credit": Decimal("0")},
                {"account": income, "debit": Decimal("0"), "credit": Decimal("500.00")},
            ],
            actor=self.admin,
            with_gst=False,
            tasks=[task],
        )
        task.refresh_from_db()
        self.assertTrue(task.bill_raised)
        self.assertEqual(task.bill_date, voucher.date)
        self.assertNotEqual(task.status, Task.STATUS_COMPLETE)

    def test_account_can_map_to_primary_group_only(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        primary = AccountGroup.objects.create(name="Sundry Debtors", parent=assets, main_head="assets")
        secondary = AccountGroup.objects.create(name="Domestic debtors", parent=primary, main_head="assets")
        form = AccountForm(
            {
                "name": "Walk-in customer",
                "primary_group": str(primary.pk),
                "secondary_group": "",
                "opening_balance": "0",
                "is_active": "on",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        acc = form.save()
        self.assertEqual(acc.group_id, primary.pk)
        self.assertEqual(acc.primary_group(), primary)
        self.assertIsNone(acc.secondary_group())
        form2 = AccountForm(
            {
                "name": "Local client",
                "primary_group": str(primary.pk),
                "secondary_group": str(secondary.pk),
                "opening_balance": "0",
                "is_active": "on",
            }
        )
        self.assertTrue(form2.is_valid(), form2.errors)
        acc2 = form2.save()
        self.assertEqual(acc2.group_id, secondary.pk)
        self.assertEqual(acc2.primary_group(), primary)
        self.assertEqual(acc2.secondary_group(), secondary)

    def test_account_master_delete(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        primary = AccountGroup.objects.create(name="Sundry Debtors", parent=assets, main_head="assets")
        acc = Account.objects.create(name="Temp ledger", group=primary)
        self.client.force_login(self.admin)
        response = self.client.post(reverse("accounts_account_delete", args=[acc.pk]), HTTP_HOST="127.0.0.1")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Account.objects.filter(pk=acc.pk).exists())
        cash = Account.objects.create(name="Cash", group=assets, system_key="cash")
        blocked = self.client.post(reverse("accounts_account_delete", args=[cash.pk]), HTTP_HOST="127.0.0.1")
        self.assertEqual(blocked.status_code, 302)
        self.assertTrue(Account.objects.filter(pk=cash.pk).exists())

    def test_account_group_delete(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        primary = AccountGroup.objects.create(name="Current Assets Del", parent=assets, main_head="assets")
        group = AccountGroup.objects.create(name="Trade receivables Del", parent=primary, main_head="assets")
        self.client.force_login(self.admin)
        blocked_primary = self.client.post(
            reverse("accounts_group_delete", args=[primary.pk]), HTTP_HOST="127.0.0.1"
        )
        self.assertEqual(blocked_primary.status_code, 302)
        self.assertTrue(AccountGroup.objects.filter(pk=primary.pk).exists())
        deleted = self.client.post(
            reverse("accounts_group_delete", args=[group.pk]), HTTP_HOST="127.0.0.1"
        )
        self.assertEqual(deleted.status_code, 302)
        self.assertFalse(AccountGroup.objects.filter(pk=group.pk).exists())
        ledger_group = AccountGroup.objects.create(name="With ledger Del", parent=primary, main_head="assets")
        Account.objects.create(name="Blocked ledger Del", group=ledger_group)
        blocked_ledger = self.client.post(
            reverse("accounts_group_delete", args=[ledger_group.pk]), HTTP_HOST="127.0.0.1"
        )
        self.assertEqual(blocked_ledger.status_code, 302)
        self.assertTrue(AccountGroup.objects.filter(pk=ledger_group.pk).exists())
        assets_blocked = self.client.post(
            reverse("accounts_group_delete", args=[assets.pk]), HTTP_HOST="127.0.0.1"
        )
        self.assertEqual(assets_blocked.status_code, 302)
        self.assertTrue(AccountGroup.objects.filter(pk=assets.pk).exists())

    def test_account_ledger_search_by_account_group_and_primary(self):
        from pathlib import Path
        from django.conf import settings

        from accounts.views import filter_account_ledgers

        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        expenses = AccountGroup.objects.get(system_key="expenses")
        debtors = AccountGroup.objects.create(name="Sundry Debtors Search", parent=assets, main_head="assets")
        domestic = AccountGroup.objects.create(name="Domestic debtors Search", parent=debtors, main_head="assets")
        office = AccountGroup.objects.create(name="Office costs Search", parent=expenses, main_head="expenses")
        buyer = Account.objects.create(name="Search Buyer Alpha", group=domestic)
        under_primary = Account.objects.create(name="Search Direct Debtor", group=debtors)
        rent = Account.objects.create(name="Search Office Rent", group=office)
        qs = Account.objects.filter(name__startswith="Search ")
        by_account = filter_account_ledgers(qs, account="buyer")
        self.assertEqual(set(by_account.values_list("pk", flat=True)), {buyer.pk})
        by_group = filter_account_ledgers(qs, account_group="domestic")
        self.assertEqual(set(by_group.values_list("pk", flat=True)), {buyer.pk})
        by_primary = filter_account_ledgers(qs, primary_group="sundry")
        self.assertEqual(set(by_primary.values_list("pk", flat=True)), {buyer.pk, under_primary.pk})
        combined = filter_account_ledgers(qs, account="rent", primary_group="office")
        self.assertEqual(set(combined.values_list("pk", flat=True)), {rent.pk})
        template = (Path(settings.BASE_DIR) / "templates/accounts/account_list.html").read_text(
            encoding="utf-8"
        )
        self.assertIn('name="account"', template)
        self.assertIn('name="account_group"', template)
        self.assertIn('name="primary_group"', template)

    def test_account_ledger_is_chart_tab_not_sidebar(self):
        from django.template.loader import render_to_string
        from django.test import RequestFactory

        from core.settings_hub import build_settings_sections

        request = RequestFactory().get("/accounts/groups/")
        request.user = self.admin
        tabs = render_to_string(
            "accounts/_chart_tabs.html",
            {"tab": "ledger", "is_accounts_admin": True},
            request=request,
        )
        self.assertIn('href="%s"' % reverse("settings_hub"), tabs)
        self.assertNotIn("Primary Group", tabs)
        self.assertNotIn("Account Ledger", tabs)
        sidebar = render_to_string(
            "includes/app_sidebar_accounts_nav.html",
            {
                "enable_accounts_module": True,
                "url_name": "accounts_account_list",
                "nav_prefix": "test",
            },
            request=request,
        )
        self.assertNotIn("Account groups", sidebar)
        self.assertNotIn("Main Account", sidebar)
        sections = build_settings_sections(self.admin)
        account = next(section for section in sections if section.title == "Account")
        self.assertEqual(
            [card.title for card in account.cards[:4]],
            ["Main Account", "Primary Group", "Account Group", "Account Ledger"],
        )
        self.assertEqual(account.cards[0].href, reverse("accounts_group_list") + "?tab=main")
        self.assertEqual(account.cards[1].href, reverse("accounts_group_list") + "?tab=parent")
        self.assertEqual(account.cards[2].href, reverse("accounts_group_list") + "?tab=account")
        self.assertEqual(account.cards[3].href, reverse("accounts_account_list"))

    def test_sales_buyer_is_sundry_debtors_only(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        expenses = AccountGroup.objects.get(system_key="expenses")
        debtors = AccountGroup.objects.create(name="Sundry Debtors", parent=assets, main_head="assets")
        sub = AccountGroup.objects.create(name="Domestic debtors", parent=debtors, main_head="assets")
        rent_group = AccountGroup.objects.create(name="Office costs", parent=expenses, main_head="expenses")
        buyer = Account.objects.create(name="Buyer one", group=debtors)
        buyer_sub = Account.objects.create(name="Buyer two", group=sub)
        other = Account.objects.create(name="Office rent", group=rent_group)
        form = InvoicePartyForm(voucher_type="sales")
        names = set(form.fields["party"].queryset.values_list("name", flat=True))
        self.assertIn(buyer.name, names)
        self.assertIn(buyer_sub.name, names)
        self.assertNotIn(other.name, names)

    def test_sales_account_tab_is_task_master_and_pure_agent_ledgers(self):
        seed_chart_of_accounts()
        income = AccountGroup.objects.get(system_key="income")
        expenses = AccountGroup.objects.get(system_key="expenses")
        master_group = AccountGroup.objects.create(
            name=self.master.name,
            parent=income,
            main_head="income",
        )
        master_ledger = Account.objects.create(name=self.master.name, group=income)
        task_named = Account.objects.create(name="GSTR-1 Apr 2026", group=master_group)
        recovery = Account.objects.create(
            name="MCA payment recovery",
            group=expenses,
            is_pure_agent=True,
        )
        rent = Account.objects.create(name="Office rent sales", group=expenses)
        income_other = Account.objects.create(name="Professional income", group=income)
        names = set(InvoiceItemForm(voucher_type="sales").fields["account"].queryset.values_list("name", flat=True))
        self.assertIn(master_ledger.name, names)
        self.assertIn(task_named.name, names)
        self.assertIn(recovery.name, names)
        self.assertNotIn(rent.name, names)
        self.assertNotIn(income_other.name, names)
        self.assertEqual(
            set(task_master_ledgers().values_list("name", flat=True)),
            {master_ledger.name, task_named.name},
        )

    def test_receipt_account_tab_is_debtors_and_capital_only(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        liabilities = AccountGroup.objects.get(system_key="liabilities")
        income = AccountGroup.objects.get(system_key="income")
        expenses = AccountGroup.objects.get(system_key="expenses")
        debtors = AccountGroup.objects.create(name="Sundry Debtors", parent=assets, main_head="assets")
        capital = AccountGroup.objects.create(name="Capital Account", parent=liabilities, main_head="liabilities")
        buyer = Account.objects.create(name="Receipt buyer", group=debtors)
        proprietor = Account.objects.create(name="Proprietor capital", group=capital)
        fees = Account.objects.create(name="Receipt fees", group=income)
        rent = Account.objects.create(name="Receipt rent", group=expenses)
        names = set(InvoiceItemForm(voucher_type="receipt").fields["account"].queryset.values_list("name", flat=True))
        self.assertIn(buyer.name, names)
        self.assertIn(proprietor.name, names)
        self.assertNotIn(fees.name, names)
        self.assertNotIn(rent.name, names)
        self.assertIn(proprietor.name, set(capital_ledgers().values_list("name", flat=True)))
        receipt_names = set(receipt_account_ledgers().values_list("name", flat=True))
        self.assertIn(buyer.name, receipt_names)
        self.assertIn(proprietor.name, receipt_names)
        self.assertNotIn(fees.name, receipt_names)
        self.assertNotIn(rent.name, receipt_names)

    def test_payment_account_tab_excludes_sundry_debtors(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        expenses = AccountGroup.objects.get(system_key="expenses")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        debtors = AccountGroup.objects.create(name="Sundry Debtors", parent=assets, main_head="assets")
        bank = Account.objects.create(name="Payment bank", group=cash_group)
        buyer = Account.objects.create(name="Payment buyer", group=debtors)
        rent = Account.objects.create(name="Payment rent", group=expenses)
        names = set(InvoiceItemForm(voucher_type="payment").fields["account"].queryset.values_list("name", flat=True))
        self.assertIn(rent.name, names)
        self.assertNotIn(buyer.name, names)
        self.assertNotIn(bank.name, names)
        self.assertIn(rent.name, set(non_cash_ledgers().values_list("name", flat=True)))
        self.assertIn(buyer.name, set(non_cash_ledgers().values_list("name", flat=True)))
        payment_names = set(payment_account_ledgers().values_list("name", flat=True))
        self.assertIn(rent.name, payment_names)
        self.assertNotIn(buyer.name, payment_names)
        self.assertNotIn(bank.name, payment_names)
        self.assertEqual(payment_names, set(line_account_ledgers("payment").values_list("name", flat=True)))

    def test_sales_invoice_print_summarises_gst_by_rate(self):
        income = _posting_ledgers()
        income.gst_rate = Decimal("18.00")
        income.hsn_code = "998314"
        income.save(update_fields=["gst_rate", "hsn_code"])
        SiteSettings.objects.update_or_create(
            pk=1,
            defaults={"company_name": "Test Firm", "gstin": "32AAAAA0000A1Z5", "address": "Trivandrum"},
        )
        self.client_row.gstin = "32ABCDE7777F1Z5"
        self.client_row.save(update_fields=["gstin"])
        buyer = sync_client_account(self.client_row)
        lines = build_invoice_lines(
            voucher_type="sales",
            party=buyer,
            items=[{"account": income, "amount": Decimal("1000.00")}],
            with_gst=True,
        )
        voucher = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 15),
            narration="Professional fee",
            lines=lines,
            actor=self.admin,
            with_gst=True,
        )
        doc = sales_invoice_document(voucher)
        self.assertEqual(doc["title"], "TAX INVOICE")
        self.assertEqual(doc["company_name"], "Test Firm")
        self.assertEqual(doc["company_address"], "Trivandrum")
        self.assertEqual(doc["company_gstin"], "32AAAAA0000A1Z5")
        self.assertEqual(len(doc["items"]), 1)
        self.assertEqual(doc["items"][0]["amount"], Decimal("1000.00"))
        self.assertEqual([row["label"] for row in doc["tax_rows"]], ["CGST 9%", "SGST 9%"])
        self.assertEqual(doc["grand_total"], Decimal("1180.00"))
        self.assertEqual(amount_in_indian_words(Decimal("1180.00")), "Rupees One Thousand One Hundred Eighty Only")

    def test_receipt_debits_cash_mode_and_credits_accounts(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        bank = Account.objects.create(name="HDFC Bank", group=cash_group)
        petty = Account.objects.create(name="Petty cash", group=cash_group)
        income = Account.objects.create(
            name="Professional income",
            group=AccountGroup.objects.get(system_key="income"),
        )
        names = set(cash_equivalent_ledgers().values_list("name", flat=True))
        self.assertIn("HDFC Bank", names)
        self.assertIn("Petty cash", names)
        self.assertNotIn("Professional income", names)
        lines = build_receipt_lines(
            mode=bank,
            items=[{"account": income, "amount": Decimal("400.00")}],
        )
        self.assertEqual(lines[0]["account"], bank)
        self.assertEqual(lines[0]["debit"], Decimal("400.00"))
        self.assertEqual(lines[1]["account"], income)
        self.assertEqual(lines[1]["credit"], Decimal("400.00"))
        voucher = post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 15),
            lines=lines,
            actor=self.admin,
        )
        self.assertTrue(voucher.number.startswith("RV/"))
        self.assertEqual(voucher.debit_total(), voucher.credit_total())

    def test_payment_credits_cash_mode_and_debits_non_cash_accounts(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        bank = Account.objects.create(name="HDFC Payment", group=cash_group)
        rent = Account.objects.create(
            name="Office rent pay",
            group=AccountGroup.objects.get(system_key="expenses"),
        )
        form = InvoicePartyForm(voucher_type="payment")
        self.assertEqual(form.fields["party"].label, "Mode of payment")
        cash_names = set(form.fields["party"].queryset.values_list("name", flat=True))
        self.assertIn("HDFC Payment", cash_names)
        self.assertNotIn("Office rent pay", cash_names)
        line_names = set(non_cash_ledgers().values_list("name", flat=True))
        self.assertIn("Office rent pay", line_names)
        self.assertNotIn("HDFC Payment", line_names)
        lines = build_payment_lines(
            mode=bank,
            items=[{"account": rent, "amount": Decimal("250.00")}],
        )
        self.assertEqual(lines[0]["account"], bank)
        self.assertEqual(lines[0]["credit"], Decimal("250.00"))
        self.assertEqual(lines[1]["account"], rent)
        self.assertEqual(lines[1]["debit"], Decimal("250.00"))
        voucher = post_voucher(
            voucher_type="payment",
            voucher_date=date(2026, 10, 16),
            lines=lines,
            actor=self.admin,
        )
        self.assertTrue(voucher.number.startswith("PV/"))
        self.assertEqual(voucher.debit_total(), voucher.credit_total())

    def test_receipt_knocks_off_sales_bills_or_posts_advance(self):
        income = _posting_ledgers()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        bank = Account.objects.create(name="Receipt bank", group=cash_group)
        buyer = sync_client_account(self.client_row)
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 10, 1),
            lines=build_invoice_lines(
                voucher_type="sales",
                party=buyer,
                items=[{"account": income, "amount": Decimal("1000.00")}],
            ),
            actor=self.admin,
        )
        bills = outstanding_bills_for_account(buyer)
        self.assertEqual(len(bills), 1)
        self.assertEqual(bills[0]["outstanding"], Decimal("1000.00"))
        self.assertEqual(bills[0]["id"], invoice.pk)

        receipt = post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 15),
            lines=build_receipt_lines(
                mode=bank,
                items=[{"account": buyer, "amount": Decimal("400.00")}],
            ),
            actor=self.admin,
            receipt_items=[
                {
                    "account": buyer,
                    "amount": Decimal("400.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": invoice.pk, "amount": "400.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_bills_for_account(buyer)[0]["outstanding"], Decimal("600.00"))
        knock = ReceiptAllocation.objects.get(receipt=receipt, is_advance=False)
        self.assertEqual(knock.against_id, invoice.pk)
        self.assertEqual(knock.amount, Decimal("400.00"))

        other = Account.objects.create(
            name="Walk-in party",
            group=AccountGroup.objects.get(system_key="assets"),
        )
        self.assertEqual(outstanding_bills_for_account(other), [])
        advance_rv = post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 16),
            lines=build_receipt_lines(
                mode=bank,
                items=[{"account": other, "amount": Decimal("250.00")}],
            ),
            actor=self.admin,
            receipt_items=[
                {
                    "account": other,
                    "amount": Decimal("250.00"),
                    "bill_alloc": {"as_advance": True, "bills": [], "advance": "250.00"},
                }
            ],
        )
        adv = ReceiptAllocation.objects.get(receipt=advance_rv)
        self.assertTrue(adv.is_advance)
        self.assertIsNone(adv.against_id)
        self.assertEqual(adv.amount, Decimal("250.00"))

        with self.assertRaises(ValidationError):
            apply_receipt_allocations(
                receipt,
                [
                    {
                        "account": buyer,
                        "amount": Decimal("10.00"),
                        "bill_alloc": {
                            "as_advance": False,
                            "bills": [{"id": invoice.pk, "amount": "700.00"}],
                        },
                    }
                ],
            )

    def test_receipt_knocks_off_opening_balance(self):
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        bank = Account.objects.create(name="OB bank", group=cash_group)
        party = Account.objects.create(
            name="Opening debtor",
            group=assets,
            opening_balance=Decimal("500.00"),
        )
        bills = outstanding_bills_for_account(party)
        self.assertEqual(len(bills), 1)
        self.assertTrue(bills[0]["is_opening"])
        self.assertEqual(bills[0]["outstanding"], Decimal("500.00"))
        receipt = post_voucher(
            voucher_type="receipt",
            voucher_date=date(2026, 10, 15),
            lines=build_receipt_lines(
                mode=bank,
                items=[{"account": party, "amount": Decimal("200.00")}],
            ),
            actor=self.admin,
            receipt_items=[
                {
                    "account": party,
                    "amount": Decimal("200.00"),
                    "bill_alloc": {
                        "as_advance": False,
                        "bills": [{"id": "opening", "amount": "200.00"}],
                        "advance": "0.00",
                    },
                }
            ],
        )
        self.assertEqual(outstanding_bills_for_account(party)[0]["outstanding"], Decimal("300.00"))
        row = ReceiptAllocation.objects.get(receipt=receipt)
        self.assertTrue(row.is_opening)
        self.assertFalse(row.is_advance)
        self.assertIsNone(row.against_id)
        self.assertEqual(row.amount, Decimal("200.00"))

    def test_account_group_and_ledger_bulk_upload(self):
        from accounts.csv_import import (
            parse_account_groups_csv,
            parse_account_ledgers_csv,
            parse_primary_groups_csv,
        )

        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        primary_csv = (
            "NAME,MAIN_HEAD,IS_ACTIVE\r\n"
            "Sundry Debtors,Assets,YES\r\n"
        ).encode("utf-8")
        primary_rows, primary_errors = parse_primary_groups_csv(primary_csv)
        self.assertEqual(primary_errors, [])
        self.assertEqual(primary_rows[0].errors, [])
        self.client.force_login(self.admin)
        session = self.client.session
        session["accounts_primary_group_import_csv"] = primary_csv.decode("utf-8")
        session.save()
        primary_confirm = self.client.post(reverse("accounts_primary_group_bulk_import"), {"confirm": "1"})
        self.assertEqual(primary_confirm.status_code, 302)
        debtors = AccountGroup.objects.get(name="Sundry Debtors", parent=assets)
        group_csv = (
            "NAME,PRIMARY_GROUP,MAIN_HEAD,IS_ACTIVE\r\n"
            "Trade receivables,Sundry Debtors,Assets,YES\r\n"
        ).encode("utf-8")
        rows, file_errors = parse_account_groups_csv(group_csv)
        self.assertEqual(file_errors, [])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].errors, [])
        self.assertEqual(rows[0].data["name"], "Trade receivables")
        self.assertEqual(rows[0].data["parent"], debtors)

        missing = parse_account_groups_csv(
            b"NAME,PRIMARY_GROUP,IS_ACTIVE\r\nWidgets,Unknown Group,YES\r\n"
        )[0]
        self.assertTrue(missing[0].errors)
        self.assertIn("PRIMARY_GROUP not found", missing[0].errors[0])

        self.client.force_login(self.admin)
        session = self.client.session
        session["accounts_group_import_csv"] = group_csv.decode("utf-8")
        session.save()
        confirm = self.client.post(reverse("accounts_group_bulk_import"), {"confirm": "1"})
        self.assertEqual(confirm.status_code, 302)
        trade = AccountGroup.objects.get(name="Trade receivables", parent=debtors)
        self.assertTrue(trade.is_account_group())

        ledger_csv = (
            "NAME,PRIMARY_GROUP,ACCOUNT_GROUP,ADDRESS,GSTIN,HSN_SAC,GST_RATE,PURE_AGENT,PAN,OPENING_BALANCE,IS_ACTIVE\r\n"
            "ABC Traders,Sundry Debtors,Trade receivables,Kochi,32AAAAA0000A1Z5,9983,18,NO,AAAAA0000A,250.50,YES\r\n"
            "Cash receipts,Sundry Debtors,,,,,,,NO,,0,YES\r\n"
        ).encode("utf-8")
        ledgers, ledger_errors = parse_account_ledgers_csv(ledger_csv)
        self.assertEqual(ledger_errors, [])
        self.assertEqual(len(ledgers), 2)
        self.assertEqual(ledgers[0].errors, [])
        self.assertEqual(ledgers[1].errors, [])
        self.assertEqual(ledgers[0].data["group"], trade)
        self.assertEqual(ledgers[1].data["group"], debtors)

        session = self.client.session
        session["accounts_ledger_import_csv"] = ledger_csv.decode("utf-8")
        session.save()
        confirm_ledgers = self.client.post(reverse("accounts_account_bulk_import"), {"confirm": "1"})
        self.assertEqual(confirm_ledgers.status_code, 302)
        abc = Account.objects.get(name="ABC Traders")
        self.assertEqual(abc.group, trade)
        self.assertEqual(abc.gstin, "32AAAAA0000A1Z5")
        self.assertEqual(abc.opening_balance, Decimal("250.50"))
        self.assertEqual(Account.objects.get(name="Cash receipts").group, debtors)

        template = self.client.get(reverse("accounts_group_bulk_import_template"))
        self.assertEqual(template.status_code, 200)
        self.assertIn(b"PRIMARY_GROUP", template.content)
        primary_template = self.client.get(reverse("accounts_primary_group_bulk_import_template"))
        self.assertEqual(primary_template.status_code, 200)
        self.assertIn(b"MAIN_HEAD", primary_template.content)
        ledger_template = self.client.get(reverse("accounts_account_bulk_import_template"))
        self.assertEqual(ledger_template.status_code, 200)
        self.assertIn(b"ACCOUNT_GROUP", ledger_template.content)


class LegacyVoucherBooksImportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="legacy-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        liab = AccountGroup.objects.get(system_key="liabilities")
        income = AccountGroup.objects.get(system_key="income")
        self.cash = Account.objects.create(name="Cash", group=assets)
        self.capital = Account.objects.create(name="Capital", group=liab)
        self.income = Account.objects.create(name="Professional income", group=income)
        self.party = Account.objects.create(name="ABC Traders", group=assets)
        self.gst_out = Account.objects.create(name="GST output", group=liab)

    def _journal_csv(self, number="JV-24-0001", debit="1000.00", credit="1000.00"):
        return (
            "DATE,VOUCHER_TYPE,NUMBER,NARRATION,ACCOUNT,DEBIT,CREDIT,SERIES\r\n"
            f"01-04-2026,Journal,{number},Opening,Cash,{debit},,JV\r\n"
            f"01-04-2026,Journal,{number},Opening,Capital,,{credit},JV\r\n"
        ).encode("utf-8")

    def test_csv_groups_two_lines_and_rejects_unbalanced(self):
        from accounts.csv_import import grouped_voucher_books, parse_voucher_books_csv

        rows, errors = parse_voucher_books_csv(self._journal_csv())
        self.assertEqual(errors, [])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].errors, [])
        self.assertEqual(rows[1].errors, [])
        books = grouped_voucher_books(rows)
        self.assertEqual(len(books), 1)
        self.assertEqual(books[0]["number"], "JV-24-0001")
        self.assertEqual(len(books[0]["lines"]), 2)

        bad, _ = parse_voucher_books_csv(self._journal_csv(debit="1000.00", credit="50.00"))
        self.assertTrue(bad[0].errors)
        self.assertTrue(any("debit must equal" in e.lower() for e in bad[0].errors))

    def test_legacy_number_does_not_consume_invoice_seq(self):
        from accounts.csv_import import parse_voucher_books_csv
        from accounts.services import allocate_voucher_number, import_legacy_vouchers

        gst_csv = (
            "DATE,VOUCHER_TYPE,NUMBER,NARRATION,ACCOUNT,DEBIT,CREDIT,SERIES\r\n"
            "02-04-2026,GST,INV/24-25/001,Fees,ABC Traders,1180.00,,GST\r\n"
            "02-04-2026,GST,INV/24-25/001,Fees,Professional income,,1000.00,GST\r\n"
            "02-04-2026,GST,INV/24-25/001,Fees,GST output,,180.00,GST\r\n"
        ).encode("utf-8")
        rows, errors = parse_voucher_books_csv(gst_csv)
        self.assertEqual(errors, [])
        self.assertTrue(all(not r.errors for r in rows))
        created = import_legacy_vouchers(rows=rows, actor=self.admin)
        self.assertEqual(created, 1)
        voucher = Voucher.objects.get(number="INV/24-25/001")
        self.assertEqual(voucher.voucher_type, "sales")
        self.assertEqual(voucher.series, SERIES_GST)
        self.assertEqual(voucher.date, date(2026, 4, 2))
        self.assertEqual(voucher.lines.count(), 3)
        self.assertEqual(InvoiceNumberSeq.objects.filter(series=SERIES_GST).count(), 0)
        next_no = allocate_voucher_number(SERIES_GST, date(2026, 4, 15))
        self.assertEqual(next_no, "GST/2026/04/00001")

    def test_existing_number_is_rejected(self):
        from accounts.csv_import import parse_voucher_books_csv

        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 4, 1),
            number="JV-24-0001",
            lines=[
                {"account": self.cash, "debit": Decimal("10.00"), "credit": Decimal("0.00")},
                {"account": self.capital, "debit": Decimal("0.00"), "credit": Decimal("10.00")},
            ],
            actor=self.admin,
        )
        rows, _ = parse_voucher_books_csv(self._journal_csv())
        self.assertTrue(any("already exists" in e for e in rows[0].errors))

    def test_same_number_allowed_across_voucher_types(self):
        from accounts.csv_import import grouped_voucher_books, parse_voucher_books_csv
        from accounts.services import import_legacy_vouchers

        csv_bytes = (
            "DATE,VOUCHER_TYPE,NUMBER,NARRATION,ACCOUNT,DEBIT,CREDIT,SERIES\r\n"
            "01-04-2026,Receipt,1,Fee received,Cash,1000.00,,RV\r\n"
            "01-04-2026,Receipt,1,Fee received,ABC Traders,,1000.00,RV\r\n"
            "02-04-2026,GST,1,Fees,ABC Traders,1180.00,,GST\r\n"
            "02-04-2026,GST,1,Fees,Professional income,,1000.00,GST\r\n"
            "02-04-2026,GST,1,Fees,GST output,,180.00,GST\r\n"
            "03-04-2026,NGST,1,Fees,ABC Traders,500.00,,NGST\r\n"
            "03-04-2026,NGST,1,Fees,Professional income,,500.00,NGST\r\n"
        ).encode("utf-8")
        rows, errors = parse_voucher_books_csv(csv_bytes)
        self.assertEqual(errors, [])
        self.assertTrue(all(not r.errors for r in rows))
        books = grouped_voucher_books(rows)
        self.assertEqual(len(books), 3)
        created = import_legacy_vouchers(rows=rows, actor=self.admin)
        self.assertEqual(created, 3)
        self.assertEqual(Voucher.objects.filter(number="1").count(), 3)
        receipt = Voucher.objects.get(number="1", voucher_type="receipt")
        gst = Voucher.objects.get(number="1", voucher_type="sales", series=SERIES_GST)
        ngst = Voucher.objects.get(number="1", voucher_type="sales", series=SERIES_NGST)
        self.assertEqual(receipt.series, "RV")
        self.assertEqual(gst.lines.count(), 3)
        self.assertEqual(ngst.lines.count(), 2)

    def test_import_gst_b2c_uses_own_series(self):
        from accounts.csv_import import grouped_voucher_books, parse_voucher_books_csv
        from accounts.services import import_legacy_vouchers

        csv_bytes = (
            "DATE,VOUCHER_TYPE,NUMBER,NARRATION,ACCOUNT,DEBIT,CREDIT,SERIES\r\n"
            "03-04-2026,GST-B2C,12,Walk-in,ABC Traders,590.00,,GSTB2C\r\n"
            "03-04-2026,GST-B2C,12,Walk-in,Professional income,,500.00,GSTB2C\r\n"
            "03-04-2026,GST-B2C,12,Walk-in,GST output,,90.00,GSTB2C\r\n"
            "03-04-2026,GST,12,B2B fees,ABC Traders,1180.00,,GST\r\n"
            "03-04-2026,GST,12,B2B fees,Professional income,,1000.00,GST\r\n"
            "03-04-2026,GST,12,B2B fees,GST output,,180.00,GST\r\n"
        ).encode("utf-8")
        rows, errors = parse_voucher_books_csv(csv_bytes)
        self.assertEqual(errors, [])
        self.assertTrue(all(not r.errors for r in rows))
        self.assertEqual(len(grouped_voucher_books(rows)), 2)
        created = import_legacy_vouchers(rows=rows, actor=self.admin)
        self.assertEqual(created, 2)
        b2c = Voucher.objects.get(number="12", voucher_type="sales", series=SERIES_GST_B2C)
        b2b = Voucher.objects.get(number="12", voucher_type="sales", series=SERIES_GST)
        self.assertEqual(b2c.date, date(2026, 4, 3))
        self.assertEqual(b2b.lines.count(), 3)
        self.assertEqual(InvoiceNumberSeq.objects.count(), 0)

    def test_same_number_same_type_is_still_rejected(self):
        from accounts.csv_import import parse_voucher_books_csv
        from accounts.services import import_legacy_vouchers

        first = (
            "DATE,VOUCHER_TYPE,NUMBER,NARRATION,ACCOUNT,DEBIT,CREDIT,SERIES\r\n"
            "01-04-2026,Receipt,1,Fee received,Cash,1000.00,,RV\r\n"
            "01-04-2026,Receipt,1,Fee received,ABC Traders,,1000.00,RV\r\n"
        ).encode("utf-8")
        rows, errors = parse_voucher_books_csv(first)
        self.assertEqual(errors, [])
        import_legacy_vouchers(rows=rows, actor=self.admin)
        again, _ = parse_voucher_books_csv(first)
        self.assertTrue(any("already exists" in e for e in again[0].errors))

    def test_confirm_import_posts_old_number(self):
        csv_bytes = self._journal_csv()
        self.client.force_login(self.admin)
        session = self.client.session
        session["accounts_voucher_books_import_csv"] = csv_bytes.decode("utf-8")
        session.save()
        confirm = self.client.post(reverse("accounts_voucher_books_bulk_import"), {"confirm": "1"})
        self.assertEqual(confirm.status_code, 302)
        voucher = Voucher.objects.get(number="JV-24-0001")
        self.assertEqual(voucher.voucher_type, "journal")
        self.assertEqual(voucher.lines.count(), 2)
        self.assertEqual(InvoiceNumberSeq.objects.count(), 0)

    def test_template_and_settings_card(self):
        from core.settings_hub import build_settings_sections

        titles = [card.title for section in build_settings_sections(self.admin) for card in section.cards]
        self.assertIn("Old books import", titles)
        self.client.force_login(self.admin)
        template = self.client.get(reverse("accounts_voucher_books_bulk_import_template"))
        self.assertEqual(template.status_code, 200)
        self.assertIn(b"VOUCHER_TYPE", template.content)
        self.assertIn(b"INV/24-25/001", template.content)

    def test_staff_cannot_import_old_books(self):
        from django.core.exceptions import PermissionDenied
        from django.test import RequestFactory

        from accounts.views import voucher_books_bulk_import

        staff = User.objects.create_user(
            email="legacy-staff@example.com",
            password="pass12345",
            is_staff=True,
        )
        request = RequestFactory().get("/accounts/vouchers/legacy-import/")
        request.user = staff
        with self.assertRaises(PermissionDenied):
            voucher_books_bulk_import(request)


class TrialBalanceReportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="tb-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        parent = AccountGroup(name="Current Assets TB", parent=assets, main_head="assets")
        parent.full_clean()
        parent.save()
        bank_group = AccountGroup(name="Bank Accounts TB", parent=parent, main_head="assets")
        bank_group.full_clean()
        bank_group.save()
        self.parent = parent
        self.bank_group = bank_group
        self.beta = Account.objects.create(
            name="Beta Bank TB",
            group=bank_group,
            opening_balance=Decimal("1000.00"),
        )
        self.alpha = Account.objects.create(
            name="Alpha Bank TB",
            group=bank_group,
            opening_balance=Decimal("200.00"),
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 3, 15),
            lines=[
                {"account": self.beta, "debit": Decimal("300.00"), "credit": Decimal("0.00")},
                {"account": self.alpha, "debit": Decimal("0.00"), "credit": Decimal("300.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": self.beta, "debit": Decimal("50.00"), "credit": Decimal("0.00")},
                {"account": self.alpha, "debit": Decimal("0.00"), "credit": Decimal("50.00")},
            ],
            actor=self.admin,
        )
        self.date_from = date(2026, 4, 1)
        self.date_to = date(2026, 10, 16)

    def test_opening_includes_pre_period_vouchers(self):
        alpha = account_totals(self.alpha, date_from=self.date_from, date_to=self.date_to)
        self.assertEqual(alpha["opening_debit"], Decimal("0.00"))
        self.assertEqual(alpha["opening_credit"], Decimal("100.00"))
        self.assertEqual(alpha["period_credit"], Decimal("50.00"))
        self.assertEqual(alpha["closing_credit"], Decimal("150.00"))
        beta = account_totals(self.beta, date_from=self.date_from, date_to=self.date_to)
        self.assertEqual(beta["opening_debit"], Decimal("1300.00"))
        self.assertEqual(beta["period_debit"], Decimal("50.00"))
        self.assertEqual(beta["closing_debit"], Decimal("1350.00"))

    def test_group_wise_parent_then_account_group_then_ledgers(self):
        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_GROUP,
            fmt=TB_FORMAT_TB,
        )
        names = [row["name"] for row in data["rows"]]
        self.assertLess(names.index("Current Assets TB"), names.index("Bank Accounts TB"))
        self.assertLess(names.index("Bank Accounts TB"), names.index("Alpha Bank TB"))
        self.assertLess(names.index("Alpha Bank TB"), names.index("Beta Bank TB"))
        parent_row = next(row for row in data["rows"] if row["name"] == "Current Assets TB")
        group_row = next(row for row in data["rows"] if row["name"] == "Bank Accounts TB")
        self.assertTrue(parent_row["is_header"])
        self.assertEqual(parent_row["kind"], "parent")
        self.assertFalse(parent_row["expandable"])
        self.assertEqual(group_row["kind"], "account_group")
        self.assertTrue(group_row["expandable"])
        self.assertEqual(parent_row["closing_debit"], Decimal("1350.00"))
        self.assertEqual(parent_row["closing_credit"], Decimal("150.00"))
        self.assertEqual(group_row["closing_debit"], Decimal("1350.00"))
        self.assertEqual(group_row["closing_credit"], Decimal("150.00"))
        alpha = next(row for row in data["rows"] if row["name"] == "Alpha Bank TB")
        self.assertTrue(alpha["collapsed"])
        self.assertEqual(alpha["parent_group_id"], self.bank_group.pk)
        self.assertFalse(data["show_group_column"])
        self.assertFalse(data["show_opening"])
        self.assertFalse(data["show_transactions"])
        from pathlib import Path
        from django.conf import settings

        template = (Path(settings.BASE_DIR) / "templates/accounts/report_trial_balance.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("js-tb-toggle", template)
        self.assertIn("bi-plus-lg", template)

    def test_drill_lists_account_groups_then_ledgers_with_both_balances(self):
        empty_group = AccountGroup(
            name="Empty Accounts TB",
            parent=self.parent,
            main_head="assets",
        )
        empty_group.full_clean()
        empty_group.save()
        Account.objects.create(name="Quiet Bank TB", group=empty_group, opening_balance=Decimal("0.00"))

        primary = trial_balance_drill(
            self.parent,
            date_from=self.date_from,
            date_to=self.date_to,
            fmt=TB_FORMAT_TB,
        )
        names = [row["name"] for row in primary["rows"]]
        self.assertEqual(primary["drill_level"], "account_groups")
        self.assertEqual(names, ["Bank Accounts TB"])
        bank = primary["rows"][0]
        self.assertEqual(bank["closing_debit"], Decimal("1350.00"))
        self.assertEqual(bank["closing_credit"], Decimal("150.00"))
        self.assertEqual(primary["totals"]["closing_debit"], Decimal("1350.00"))
        self.assertEqual(primary["totals"]["closing_credit"], Decimal("150.00"))

        ledgers = trial_balance_drill(
            self.bank_group,
            date_from=self.date_from,
            date_to=self.date_to,
            fmt=TB_FORMAT_TB,
        )
        self.assertEqual(ledgers["drill_level"], "ledgers")
        self.assertEqual([row["name"] for row in ledgers["rows"]], ["Alpha Bank TB", "Beta Bank TB"])
        self.assertEqual(ledgers["rows"][0]["closing_credit"], Decimal("150.00"))
        self.assertEqual(ledgers["rows"][1]["closing_debit"], Decimal("1350.00"))

        assets = AccountGroup.objects.get(system_key="assets")
        heads = trial_balance_drill(assets, date_from=self.date_from, date_to=self.date_to, fmt=TB_FORMAT_TB)
        self.assertEqual(heads["drill_level"], "primary_groups")
        self.assertIn("Current Assets TB", [row["name"] for row in heads["rows"]])

    def test_group_drill_pages_open_the_next_level_for_the_same_period(self):
        from accounts.views import _annotate_tb_open_urls

        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_GROUP,
            fmt=TB_FORMAT_TB,
        )
        _annotate_tb_open_urls(data["rows"], self.date_from, self.date_to, TB_REPORT_GROUP, TB_FORMAT_TB)
        alpha = next(row for row in data["rows"] if row["name"] == "Alpha Bank TB")
        bank = next(row for row in data["rows"] if row["name"] == "Bank Accounts TB")
        primary_row = next(row for row in data["rows"] if row["name"] == "Current Assets TB")
        self.assertIn(f"account={self.alpha.pk}", alpha["open_url"])
        self.assertIn("date_from=2026-04-01", alpha["open_url"])
        self.assertIn("date_to=2026-10-16", alpha["open_url"])
        self.assertIn(reverse("accounts_ledger"), alpha["open_url"])
        group_path = reverse("accounts_trial_balance_group", args=[self.bank_group.pk])
        primary_path = reverse("accounts_trial_balance_group", args=[self.parent.pk])
        self.assertTrue(bank["open_url"].startswith(group_path))
        self.assertTrue(primary_row["open_url"].startswith(primary_path))

        primary = trial_balance_drill(self.parent, date_from=self.date_from, date_to=self.date_to, fmt=TB_FORMAT_TB)
        _annotate_tb_open_urls(primary["rows"], self.date_from, self.date_to, TB_REPORT_GROUP, TB_FORMAT_TB)
        self.assertTrue(primary["rows"][0]["open_url"].startswith(group_path))

        ledgers = trial_balance_drill(
            self.bank_group, date_from=self.date_from, date_to=self.date_to, fmt=TB_FORMAT_TB
        )
        _annotate_tb_open_urls(ledgers["rows"], self.date_from, self.date_to, TB_REPORT_GROUP, TB_FORMAT_TB)
        opened = {row["name"]: row["open_url"] for row in ledgers["rows"]}
        self.assertIn(f"account={self.alpha.pk}", opened["Alpha Bank TB"])
        self.assertIn(f"account={self.beta.pk}", opened["Beta Bank TB"])
        self.assertNotIn("tb_report", opened["Alpha Bank TB"])

    def test_group_wise_heads_are_liabilities_assets_income_expenses(self):
        income = AccountGroup.objects.get(system_key="income")
        expenses = AccountGroup.objects.get(system_key="expenses")
        liabilities = AccountGroup.objects.get(system_key="liabilities")
        Account.objects.create(
            name="TB Fees Inc",
            group=income,
            opening_balance=Decimal("10.00"),
        )
        Account.objects.create(
            name="TB Rent Exp",
            group=expenses,
            opening_balance=Decimal("5.00"),
        )
        Account.objects.create(
            name="TB Capital",
            group=liabilities,
            opening_balance=Decimal("15.00"),
        )
        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_GROUP,
            fmt=TB_FORMAT_TB,
        )
        heads = [row["name"] for row in data["rows"] if row["indent"] == 0]
        self.assertEqual(heads[:4], ["Liabilities", "Assets", "Income", "Expenses"])

    def test_ledger_wise_is_alphabetical_with_group_column(self):
        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_LEDGER,
            fmt=TB_FORMAT_TRANSACTIONS,
        )
        ledgers = [row for row in data["rows"] if row["kind"] == "ledger"]
        names = [row["name"] for row in ledgers]
        self.assertEqual(names, sorted(names, key=str.lower))
        self.assertTrue(data["show_group_column"])
        self.assertTrue(data["show_opening"])
        self.assertTrue(data["show_transactions"])
        alpha = next(row for row in ledgers if row["name"] == "Alpha Bank TB")
        self.assertEqual(alpha["group"].name, "Bank Accounts TB")

    def test_opening_format_hides_period_columns(self):
        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_GROUP,
            fmt=TB_FORMAT_OPENING,
        )
        self.assertTrue(data["show_opening"])
        self.assertFalse(data["show_transactions"])

    def test_excel_and_pdf_export(self):
        from io import BytesIO
        from openpyxl import load_workbook

        data = trial_balance(
            date_from=self.date_from,
            date_to=self.date_to,
            report=TB_REPORT_LEDGER,
            fmt=TB_FORMAT_TRANSACTIONS,
        )
        xlsx = trial_balance_xlsx_bytes(
            data, date_from=self.date_from, date_to=self.date_to, company_name="Test Firm"
        )
        wb = load_workbook(BytesIO(xlsx))
        values = [cell.value for row in wb.active.iter_rows() for cell in row]
        self.assertIn("Alpha Bank TB", values)
        self.assertIn("Group", values)
        pdf = trial_balance_pdf_bytes(
            data, date_from=self.date_from, date_to=self.date_to, company_name="Test Firm"
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Alpha Bank TB", pdf)
        _assert_pdf_structure(self, pdf)

        self.client.force_login(self.admin)
        xlsx_res = self.client.get(
            reverse("accounts_trial_balance"),
            {
                "date_from": self.date_from,
                "date_to": self.date_to,
                "tb_report": TB_REPORT_GROUP,
                "report_format": TB_FORMAT_TB,
                "export": "xlsx",
            },
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_res["Content-Type"])
        pdf_res = self.client.get(
            reverse("accounts_trial_balance"),
            {
                "date_from": self.date_from,
                "date_to": self.date_to,
                "tb_report": TB_REPORT_GROUP,
                "report_format": TB_FORMAT_TB,
                "export": "pdf",
            },
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")
        _assert_pdf_structure(self, pdf_res.content)
        self.assertIn(b"/F2 12.0 Tf", pdf_res.content)
        self.assertNotIn(b"0.91 0.92 0.93 rg", pdf_res.content)


def _pnl_group(name, parent):
    group = AccountGroup(name=name, parent=parent, main_head=parent.main_head)
    group.full_clean()
    group.save()
    return group


class ProfitAndLossFormatTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="pnl-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        expenses = AccountGroup.objects.get(system_key="expenses")
        income = AccountGroup.objects.get(system_key="income")
        cash_group = _pnl_group("Cash and Cash Equivalents", assets)
        stock_group = _pnl_group("Stock-in-hand", assets)
        direct_exp = _pnl_group("Direct Expenses", expenses)
        indirect_exp = _pnl_group("Indirect Expenses", expenses)
        depreciation = _pnl_group("Depreciation and Amortization", expenses)
        finance = _pnl_group("Finance Cost", expenses)
        purchase_group = _pnl_group("Purchase", expenses)
        direct_inc = _pnl_group("Direct Income", income)
        indirect_inc = _pnl_group("Indirect Income", income)
        self.cash = Account.objects.create(name="P&L cash", group=cash_group)
        self.stock = Account.objects.create(
            name="Stock",
            group=stock_group,
            opening_balance=Decimal("2500.00"),
        )
        purchase = Account.objects.create(name="Purchase", group=purchase_group)
        freight = Account.objects.create(name="Freight", group=direct_exp)
        customs = Account.objects.create(name="Customs Duty", group=direct_exp)
        sales = Account.objects.create(name="Sales", group=direct_inc)
        advert = Account.objects.create(name="Advertisement", group=indirect_exp)
        salary = Account.objects.create(name="Salary", group=indirect_exp)
        dep = Account.objects.create(name="Depreciation", group=depreciation)
        interest = Account.objects.create(name="Interest", group=finance)
        interest_inc = Account.objects.create(name="Interest Income", group=indirect_inc)
        scrap = Account.objects.create(name="Scrap Sales", group=indirect_inc)
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": purchase, "debit": Decimal("10000.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("10000.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": freight, "debit": Decimal("1000.00"), "credit": Decimal("0.00")},
                {"account": customs, "debit": Decimal("500.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("1500.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": self.cash, "debit": Decimal("10000.00"), "credit": Decimal("0.00")},
                {"account": sales, "debit": Decimal("0.00"), "credit": Decimal("10000.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": self.stock, "debit": Decimal("2500.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("2500.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": advert, "debit": Decimal("250.00"), "credit": Decimal("0.00")},
                {"account": salary, "debit": Decimal("300.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("550.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": dep, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("100.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": interest, "debit": Decimal("10.00"), "credit": Decimal("0.00")},
                {"account": self.cash, "debit": Decimal("0.00"), "credit": Decimal("10.00")},
            ],
            actor=self.admin,
        )
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": self.cash, "debit": Decimal("500.00"), "credit": Decimal("0.00")},
                {"account": interest_inc, "debit": Decimal("0.00"), "credit": Decimal("200.00")},
                {"account": scrap, "debit": Decimal("0.00"), "credit": Decimal("300.00")},
            ],
            actor=self.admin,
        )

    def test_horizontal_pnl_matches_trading_format(self):
        data = profit_and_loss(date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        self.assertEqual(data["gross_profit"], Decimal("1000.00"))
        self.assertEqual(data["net_profit"], Decimal("840.00"))
        trading_left = [row["left"]["name"] for row in data["trading_rows"]]
        trading_right = [row["right"]["name"] for row in data["trading_rows"]]
        self.assertIn("Opening Stock", trading_left)
        self.assertIn("Purchase", trading_left)
        self.assertIn("Direct Expenses", trading_left)
        self.assertIn("Gross Profit C/f", trading_left)
        self.assertIn("Direct Income", trading_right)
        self.assertNotIn("Sales", trading_right)
        self.assertIn("Closing Stock", trading_right)
        direct_income = next(row for row in data["trading_rows"] if row["right"]["name"] == "Direct Income")
        self.assertTrue(direct_income["right"]["expandable"])
        self.assertEqual(direct_income["right"]["outer"], Decimal("10000.00"))
        self.assertEqual(
            [child["name"] for child in direct_income["right"]["children"]],
            ["Sales"],
        )
        sales = Account.objects.get(name="Sales")
        self.assertEqual(direct_income["right"]["children"][0]["account_id"], sales.pk)
        freight_row = next(row for row in data["trading_rows"] if row["left"]["name"] == "Freight")
        self.assertEqual(freight_row["left"]["account_id"], Account.objects.get(name="Freight").pk)
        self.assertIsNone(direct_income["right"]["account_id"])
        gp = next(row for row in data["trading_rows"] if row["left"]["name"] == "Gross Profit C/f")
        self.assertEqual(gp["left"]["outer"], Decimal("1000.00"))
        totals = data["trading_rows"][-1]
        self.assertEqual(totals["left"]["name"], "Total")
        self.assertEqual(totals["left"]["outer"], Decimal("15000.00"))
        self.assertEqual(totals["right"]["outer"], Decimal("15000.00"))
        pnl_left = [row["left"]["name"] for row in data["pnl_rows"]]
        pnl_right = [row["right"]["name"] for row in data["pnl_rows"]]
        self.assertIn("Indirect Expenses", pnl_left)
        self.assertIn("Net Profit/Loss", pnl_left)
        self.assertIn("Gross Profit B/f", pnl_right)
        self.assertIn("Indirect Income", pnl_right)
        net = next(row for row in data["pnl_rows"] if row["left"]["name"] == "Net Profit/Loss")
        self.assertEqual(net["left"]["outer"], Decimal("840.00"))
        pnl_totals = data["pnl_rows"][-1]
        self.assertEqual(pnl_totals["left"]["outer"], Decimal("1500.00"))
        self.assertEqual(pnl_totals["right"]["outer"], Decimal("1500.00"))
        direct_ledgers = [
            row for row in data["trading_rows"] if row["left"]["name"] in ("Freight", "Customs Duty")
        ]
        self.assertEqual(
            {row["left"]["name"]: row["left"]["inner"] for row in direct_ledgers},
            {"Freight": Decimal("1000.00"), "Customs Duty": Decimal("500.00")},
        )
        self.assertEqual(direct_ledgers[-1]["left"]["outer"], Decimal("1500.00"))

    def test_direct_income_groups_at_account_group_level(self):
        income = AccountGroup.objects.get(system_key="income")
        direct_inc = AccountGroup.objects.get(name="Direct Income", parent=income)
        audit = _pnl_group("Audit & Assurance", direct_inc)
        accounting = _pnl_group("Accounting Services", direct_inc)
        concurrent = Account.objects.create(name="Concurrent Audit", group=audit)
        internal = Account.objects.create(name="Internal Audit", group=audit)
        monthly = Account.objects.create(name="Accounting Services - Monthly", group=accounting)
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 6, 1),
            lines=[
                {"account": self.cash, "debit": Decimal("473000.00"), "credit": Decimal("0.00")},
                {"account": concurrent, "debit": Decimal("0.00"), "credit": Decimal("107000.00")},
                {"account": internal, "debit": Decimal("0.00"), "credit": Decimal("296000.00")},
                {"account": monthly, "debit": Decimal("0.00"), "credit": Decimal("70000.00")},
            ],
            actor=self.admin,
        )
        data = profit_and_loss(date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        trading_right = [row["right"]["name"] for row in data["trading_rows"]]
        self.assertIn("Accounting Services", trading_right)
        self.assertIn("Audit & Assurance", trading_right)
        self.assertIn("Direct Income", trading_right)
        self.assertNotIn("Concurrent Audit", trading_right)
        self.assertNotIn("Internal Audit", trading_right)
        self.assertNotIn("Accounting Services - Monthly", trading_right)
        audit_row = next(row for row in data["trading_rows"] if row["right"]["name"] == "Audit & Assurance")
        self.assertTrue(audit_row["right"]["expandable"])
        self.assertEqual(audit_row["right"]["outer"], Decimal("403000.00"))
        self.assertEqual(
            [child["name"] for child in audit_row["right"]["children"]],
            ["Concurrent Audit", "Internal Audit"],
        )
        accounting_row = next(
            row for row in data["trading_rows"] if row["right"]["name"] == "Accounting Services"
        )
        self.assertEqual(accounting_row["right"]["outer"], Decimal("70000.00"))
        self.assertEqual(
            [child["name"] for child in accounting_row["right"]["children"]],
            ["Accounting Services - Monthly"],
        )
        self.assertEqual(
            [child["account_id"] for child in audit_row["right"]["children"]],
            [concurrent.pk, internal.pk],
        )
        self.assertEqual(accounting_row["right"]["children"][0]["account_id"], monthly.pk)

    def test_livestock_debtor_is_not_closing_stock(self):
        from accounts.services import ensure_sundry_debtors_group, profit_and_loss

        debtor = Account.objects.create(
            name="MAZHAITHULI LIVESTOCK FARMER PRODUCER COMPANY LIMITED",
            group=ensure_sundry_debtors_group(),
            opening_balance=Decimal("37540.00"),
        )
        data = profit_and_loss(date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        trading_right = [row["right"]["name"] for row in data["trading_rows"]]
        self.assertIn("Closing Stock", trading_right)
        closing = next(row for row in data["trading_rows"] if row["right"]["name"] == "Closing Stock")
        self.assertEqual(closing["right"]["outer"], Decimal("5000.00"))
        self.assertNotEqual(closing["right"]["outer"], debtor.opening_balance)

    def test_excel_and_pdf_export(self):
        from io import BytesIO
        from openpyxl import load_workbook

        data = profit_and_loss(date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        xlsx = profit_and_loss_xlsx_bytes(
            data,
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        values = [cell.value for row in load_workbook(BytesIO(xlsx)).active.iter_rows() for cell in row]
        self.assertTrue(any(isinstance(value, str) and value.strip() == "Sales" for value in values))
        self.assertIn("Gross Profit C/f", values)
        pdf = profit_and_loss_pdf_bytes(
            data,
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Sales", pdf)
        _assert_pdf_structure(self, pdf)
        self.client.force_login(self.admin)
        xlsx_res = self.client.get(
            reverse("accounts_profit_loss"),
            {"date_from": "2026-04-01", "date_to": "2026-10-16", "export": "xlsx"},
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_res["Content-Type"])
        pdf_res = self.client.get(
            reverse("accounts_profit_loss"),
            {"date_from": "2026-04-01", "date_to": "2026-10-16", "export": "pdf"},
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")
        _assert_pdf_structure(self, pdf_res.content)


class BalanceSheetFormatTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="bs-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        liabilities = AccountGroup.objects.get(system_key="liabilities")
        fixed = _pnl_group("Fixed Assets", assets)
        non_cur_assets = _pnl_group("Non Current Assets", assets)
        current_assets = _pnl_group("Current Assets", assets)
        bank_group = _pnl_group("Bank Accounts", current_assets)
        capital = _pnl_group("Capital Account", liabilities)
        non_cur_liab = _pnl_group("Non Current Liabilities", liabilities)
        current_liab = _pnl_group("Current Liabilities", liabilities)
        creditors = _pnl_group("Sundry Creditors", current_liab)
        Account.objects.create(name="Furniture", group=fixed, opening_balance=Decimal("5000.00"))
        Account.objects.create(
            name="Investments", group=non_cur_assets, opening_balance=Decimal("2000.00")
        )
        Account.objects.create(
            name="Till cash BS", group=current_assets, opening_balance=Decimal("800.00")
        )
        Account.objects.create(
            name="HDFC Bank BS", group=bank_group, opening_balance=Decimal("1200.00")
        )
        Account.objects.create(
            name="Owner capital", group=capital, opening_balance=Decimal("7000.00")
        )
        Account.objects.create(
            name="Term loan", group=non_cur_liab, opening_balance=Decimal("1500.00")
        )
        Account.objects.create(
            name="GST payable", group=current_liab, opening_balance=Decimal("300.00")
        )
        Account.objects.create(
            name="Vendor A", group=creditors, opening_balance=Decimal("200.00")
        )
        income = AccountGroup.objects.get(system_key="income")
        direct_inc = _pnl_group("Direct Income", income)
        fees = Account.objects.create(name="Filing fees BS", group=direct_inc)
        cash = Account.objects.get(name="Till cash BS")
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": cash, "debit": Decimal("500.00"), "credit": Decimal("0.00")},
                {"account": fees, "debit": Decimal("0.00"), "credit": Decimal("500.00")},
            ],
            actor=self.admin,
        )

    def test_liabilities_left_assets_right_with_group_order(self):
        data = balance_sheet(date_to=date(2026, 10, 16))
        left = [row["left"]["name"] for row in data["statement_rows"] if row["left"]["name"]]
        right = [row["right"]["name"] for row in data["statement_rows"] if row["right"]["name"]]
        self.assertLess(right.index("Fixed Assets"), right.index("Non Current Assets"))
        self.assertLess(right.index("Non Current Assets"), right.index("Current Assets"))
        self.assertLess(right.index("Current Assets"), right.index("Bank Accounts"))
        self.assertLess(right.index("Bank Accounts"), right.index("HDFC Bank BS"))
        self.assertLess(right.index("Current Assets"), right.index("Till cash BS"))
        self.assertEqual(right[right.index("Fixed Assets") + 1], "Furniture")
        self.assertLess(left.index("Capital Account"), left.index("Owner capital"))
        self.assertLess(left.index("Owner capital"), left.index("Net Profit"))
        self.assertLess(left.index("Net Profit"), left.index("Non Current Liabilities"))
        self.assertLess(left.index("Non Current Liabilities"), left.index("Current Liabilities"))
        self.assertLess(left.index("Current Liabilities"), left.index("Sundry Creditors"))
        self.assertLess(left.index("Sundry Creditors"), left.index("Vendor A"))
        self.assertIn("GST payable", left)
        np_row = next(row for row in data["statement_rows"] if row["left"]["name"] == "Net Profit")
        self.assertEqual(np_row["left"]["indent"], 1)
        self.assertEqual(np_row["left"]["inner"], Decimal("500.00"))
        cash_row = next(row for row in data["statement_rows"] if row["right"]["name"] == "Till cash BS")
        self.assertEqual(cash_row["right"]["indent"], 1)
        self.assertEqual(cash_row["right"]["account_id"], Account.objects.get(name="Till cash BS").pk)
        self.assertIsNone(np_row["left"]["account_id"])
        furniture_row = next(row for row in data["statement_rows"] if row["right"]["name"] == "Furniture")
        self.assertEqual(furniture_row["right"]["account_id"], Account.objects.get(name="Furniture").pk)
        hdfc_row = next(row for row in data["statement_rows"] if row["right"]["name"] == "HDFC Bank BS")
        self.assertEqual(hdfc_row["right"]["indent"], 2)
        bank_row = next(row for row in data["statement_rows"] if row["right"]["name"] == "Bank Accounts")
        self.assertEqual(bank_row["right"]["kind"], "subheader")
        self.assertEqual(bank_row["right"]["indent"], 1)
        totals = data["statement_rows"][-1]
        self.assertEqual(totals["left"]["name"], "Total")
        self.assertEqual(totals["left"]["outer"], totals["right"]["outer"])
        self.assertEqual(totals["left"]["outer"], Decimal("9500.00"))

    def test_excel_and_pdf_export(self):
        from io import BytesIO
        from openpyxl import load_workbook

        data = balance_sheet(date_to=date(2026, 10, 16))
        xlsx = balance_sheet_xlsx_bytes(
            data,
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        values = [cell.value for row in load_workbook(BytesIO(xlsx)).active.iter_rows() for cell in row]
        names = [str(value).strip() for value in values if value not in (None, "")]
        self.assertIn("Liabilities", names)
        self.assertIn("Assets", names)
        self.assertIn("Capital Account", names)
        self.assertIn("Fixed Assets", names)
        self.assertIn("Net Profit", names)
        pdf = balance_sheet_pdf_bytes(
            data,
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Liabilities", pdf)
        self.assertIn(b"Assets", pdf)
        self.assertIn(b"Capital Account", pdf)
        _assert_pdf_structure(self, pdf)
        self.client.force_login(self.admin)
        xlsx_res = self.client.get(
            reverse("accounts_balance_sheet"),
            {"date_to": "2026-10-16", "export": "xlsx"},
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_res["Content-Type"])
        self.assertIn("balance-sheet-2026-10-16.xlsx", xlsx_res["Content-Disposition"])
        pdf_res = self.client.get(
            reverse("accounts_balance_sheet"),
            {"date_to": "2026-10-16", "export": "pdf"},
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")
        self.assertIn("balance-sheet-2026-10-16.pdf", pdf_res["Content-Disposition"])
        _assert_pdf_structure(self, pdf_res.content)


class LedgerFormatTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="ledger-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = _pnl_group("Cash and Cash Equivalents", assets)
        debtors = _pnl_group("Sundry Debtors", assets)
        self.cash = Account.objects.create(name="Cash", group=cash_group)
        self.party = Account.objects.create(name="Antony Edgar Ledger", group=debtors)
        self.voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 1),
            narration="Receipt from party",
            lines=[
                {"account": self.cash, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": self.party, "debit": Decimal("0.00"), "credit": Decimal("100.00")},
            ],
            actor=self.admin,
        )

    def test_party_ledger_shows_cash_as_counterpart(self):
        rows = ledger_lines(self.party, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        entry = next(row for row in rows if row.get("voucher") and row["voucher"].pk == self.voucher.pk)
        self.assertEqual(entry["date"], date(2026, 10, 1))
        self.assertEqual(entry["number"], self.voucher.number)
        self.assertEqual(entry["ledger_account"], "Cash")
        self.assertEqual(entry["narration"], "Receipt from party")
        self.assertEqual(entry["credit"], Decimal("100.00"))
        self.assertEqual(entry["debit"], Decimal("0.00"))
        self.assertEqual(entry["amount"], Decimal("100.00"))
        closing = next(row for row in rows if row["kind"] == "closing")
        total = next(row for row in rows if row["kind"] == "total")
        self.assertEqual(closing["ledger_account"], "Closing Balance")
        self.assertEqual(closing["debit"], Decimal("100.00"))
        self.assertEqual(total["debit"], total["credit"])
        self.assertEqual(total["debit"], Decimal("100.00"))

    def test_cash_ledger_shows_party_as_counterpart(self):
        rows = ledger_lines(self.cash, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        entry = next(row for row in rows if row.get("voucher") and row["voucher"].pk == self.voucher.pk)
        self.assertEqual(entry["ledger_account"], "Antony Edgar Ledger")
        self.assertEqual(entry["debit"], Decimal("100.00"))
        self.assertEqual(entry["credit"], Decimal("0.00"))

    def test_compound_voucher_shows_this_ledger_amount_and_opposite_party(self):
        bank = Account.objects.create(name="HDFC Ledger", group=self.cash.group)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 2),
            narration="Split receipt",
            lines=[
                {"account": self.party, "debit": Decimal("0.00"), "credit": Decimal("150.00")},
                {"account": self.cash, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": bank, "debit": Decimal("50.00"), "credit": Decimal("0.00")},
            ],
            actor=self.admin,
        )
        rows = [
            row
            for row in ledger_lines(self.party, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
            if row.get("voucher") and row["voucher"].pk == voucher.pk
        ]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["ledger_account"], "Cash")
        self.assertEqual(rows[0]["credit"], Decimal("150.00"))
        self.assertEqual(rows[0]["debit"], Decimal("0.00"))
        cash_rows = [
            row
            for row in ledger_lines(self.cash, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
            if row.get("voucher") and row["voucher"].pk == voucher.pk
        ]
        self.assertEqual(len(cash_rows), 1)
        self.assertEqual(cash_rows[0]["ledger_account"], "Antony Edgar Ledger")
        self.assertEqual(cash_rows[0]["debit"], Decimal("100.00"))

    def test_gst_sales_ledger_hides_same_side_tax_lines(self):
        income = AccountGroup.objects.get(system_key="income")
        liabilities = AccountGroup.objects.get(system_key="liabilities")
        fees = Account.objects.create(name="GST Monthly Fees", group=_pnl_group("Direct Income", income))
        cgst = Account.objects.create(name="CGST Output 9%", group=_pnl_group("Duties and Taxes", liabilities))
        sgst = Account.objects.create(name="SGST Output 9%", group=cgst.group)
        voucher = post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 10, 7),
            lines=[
                {"account": self.party, "debit": Decimal("1180.00"), "credit": Decimal("0.00")},
                {"account": fees, "debit": Decimal("0.00"), "credit": Decimal("1000.00")},
                {"account": cgst, "debit": Decimal("0.00"), "credit": Decimal("90.00")},
                {"account": sgst, "debit": Decimal("0.00"), "credit": Decimal("90.00")},
            ],
            actor=self.admin,
        )
        fee_rows = [
            row
            for row in ledger_lines(fees, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
            if row.get("voucher") and row["voucher"].pk == voucher.pk
        ]
        self.assertEqual(len(fee_rows), 1)
        self.assertEqual(fee_rows[0]["ledger_account"], "Antony Edgar Ledger")
        self.assertEqual(fee_rows[0]["credit"], Decimal("1000.00"))
        self.assertEqual(fee_rows[0]["debit"], Decimal("0.00"))
        names = [row["ledger_account"] for row in fee_rows]
        self.assertNotIn("CGST Output 9%", names)
        self.assertNotIn("SGST Output 9%", names)
        party_rows = [
            row
            for row in ledger_lines(self.party, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
            if row.get("voucher") and row["voucher"].pk == voucher.pk
        ]
        self.assertEqual(len(party_rows), 1)
        self.assertEqual(party_rows[0]["ledger_account"], "GST Monthly Fees")
        self.assertEqual(party_rows[0]["debit"], Decimal("1180.00"))
        cgst_rows = [
            row
            for row in ledger_lines(cgst, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
            if row.get("voucher") and row["voucher"].pk == voucher.pk
        ]
        self.assertEqual(cgst_rows[0]["ledger_account"], "Antony Edgar Ledger")
        self.assertEqual(cgst_rows[0]["credit"], Decimal("90.00"))

    def test_excel_and_pdf_export(self):
        from io import BytesIO
        from openpyxl import load_workbook

        rows = ledger_lines(self.party, date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        xlsx = ledger_xlsx_bytes(
            rows,
            account_name=self.party.name,
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        values = [str(cell.value).strip() for row in load_workbook(BytesIO(xlsx)).active.iter_rows() for cell in row if cell.value not in (None, "")]
        self.assertIn("Cash", values)
        self.assertIn("Closing Balance", values)
        self.assertIn("Total", values)
        pdf = ledger_pdf_bytes(
            rows,
            account_name=self.party.name,
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
        )
        self.assertTrue(pdf.startswith(b"%PDF"))
        self.assertIn(b"Cash", pdf)
        self.assertIn(b"Closing Balance", pdf)
        self.client.force_login(self.admin)
        xlsx_res = self.client.get(
            reverse("accounts_ledger"),
            {"account": str(self.party.pk), "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "xlsx"},
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("spreadsheetml", xlsx_res["Content-Type"])
        pdf_res = self.client.get(
            reverse("accounts_ledger"),
            {"account": str(self.party.pk), "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "pdf"},
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")

    def test_voucher_back_link_returns_to_ledger(self):
        from accounts.views import _ledger_back_url

        ledger_url = f"{reverse('accounts_ledger')}?account={self.party.pk}"
        self.assertEqual(_ledger_back_url(ledger_url), ledger_url)
        self.assertEqual(_ledger_back_url("/accounts/vouchers/"), "")
        self.assertEqual(_ledger_back_url("https://example.com/accounts/reports/ledger/"), "")

    def test_report_return_url_allows_pnl_and_balance_sheet(self):
        from accounts.views import _report_return_url

        pnl = f"{reverse('accounts_profit_loss')}?date_from=2026-04-01&date_to=2026-10-16"
        bs = f"{reverse('accounts_balance_sheet')}?date_to=2026-10-16"
        tb = f"{reverse('accounts_trial_balance')}?date_from=2026-04-01&date_to=2026-10-16"
        group = f"{reverse('accounts_trial_balance_group', args=[self.party.group_id])}?date_from=2026-04-01"
        self.assertEqual(_report_return_url(pnl), pnl)
        self.assertEqual(_report_return_url(bs), bs)
        self.assertEqual(_report_return_url(tb), tb)
        self.assertEqual(_report_return_url(group), group)
        self.assertEqual(_report_return_url("/accounts/vouchers/"), "")
        self.assertEqual(_report_return_url("https://example.com/accounts/reports/profit-loss/"), "")

    def test_all_ledger_groups_by_parent_and_exports(self):
        from io import BytesIO
        from openpyxl import load_workbook

        from accounts.export import ledger_xlsx_bytes

        sections = all_ledger_sections(date_from=date(2026, 4, 1), date_to=date(2026, 10, 16))
        names = [section["account"].name for section in sections]
        self.assertIn("Cash", names)
        self.assertIn("Antony Edgar Ledger", names)
        cash = next(section for section in sections if section["account"].name == "Cash")
        party = next(section for section in sections if section["account"].name == "Antony Edgar Ledger")
        self.assertEqual(cash["parent_group"], "Cash and Cash Equivalents")
        self.assertEqual(party["parent_group"], "Sundry Debtors")
        self.assertLess(names.index("Cash"), names.index("Antony Edgar Ledger"))
        xlsx = ledger_xlsx_bytes(
            [],
            date_from=date(2026, 4, 1),
            date_to=date(2026, 10, 16),
            company_name="Test Firm",
            sections=sections,
        )
        values = [
            str(cell.value).strip()
            for row in load_workbook(BytesIO(xlsx)).active.iter_rows()
            for cell in row
            if cell.value not in (None, "")
        ]
        self.assertIn("All Ledger", values)
        self.assertIn("Cash and Cash Equivalents", values)
        self.assertIn("Sundry Debtors", values)
        self.client.force_login(self.admin)
        xlsx_res = self.client.get(
            reverse("accounts_ledger"),
            {"account": "all", "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "xlsx"},
        )
        self.assertEqual(xlsx_res.status_code, 200)
        self.assertIn("ledger-all-", xlsx_res["Content-Disposition"])
        pdf_res = self.client.get(
            reverse("accounts_ledger"),
            {"account": "all", "date_from": "2026-04-01", "date_to": "2026-10-16", "export": "pdf"},
        )
        self.assertEqual(pdf_res.status_code, 200)
        self.assertEqual(pdf_res["Content-Type"], "application/pdf")
        _assert_pdf_structure(self, pdf_res.content)


class PdfBuilderTests(SimpleTestCase):
    def test_contents_and_kids_point_at_real_objects(self):
        pdf = _build_pdf([[_pdf_text(72, 720, "Hello")]], 595, 842)
        _assert_pdf_structure(self, pdf)
        self.assertIn(b"/Contents 5 0 R", pdf)
        self.assertIn(b"/Kids [6 0 R]", pdf)
        self.assertIn(b"5 0 obj", pdf)
        self.assertIn(b"6 0 obj", pdf)
        self.assertNotIn(b"/Kids [7 0 R]", pdf)


class ResetTestDataAccountsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="reset-acc-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        income = _posting_ledgers()
        assets = AccountGroup.objects.get(system_key="assets")
        bank_parent = AccountGroup.objects.create(
            name="Bank Accounts Reset",
            parent=assets,
            main_head="assets",
        )
        self.cash = Account.objects.create(name="Reset cash", group=bank_parent)
        post_voucher(
            voucher_type="journal",
            voucher_date=date(2026, 5, 1),
            lines=[
                {"account": self.cash, "debit": Decimal("100.00"), "credit": Decimal("0.00")},
                {"account": income, "debit": Decimal("0.00"), "credit": Decimal("100.00")},
            ],
            actor=self.admin,
        )

    def test_accounts_wipe_removes_vouchers_and_user_groups(self):
        from core.reset_data import WipeOptions, count_local_data, wipe_local_data

        counts = count_local_data()
        self.assertGreater(counts["accounts_vouchers"], 0)
        self.assertGreater(counts["accounts_ledgers"], 0)
        self.assertGreater(counts["accounts_groups"], 0)
        self.assertGreater(InvoiceNumberSeq.objects.count(), 0)

        deleted = wipe_local_data(WipeOptions(accounts=True))
        self.assertGreater(deleted["accounts_vouchers"], 0)
        self.assertEqual(Voucher.objects.count(), 0)
        self.assertEqual(ReceiptAllocation.objects.count(), 0)
        self.assertEqual(InvoiceNumberSeq.objects.count(), 0)
        self.assertEqual(Account.objects.count(), 0)
        self.assertEqual(AccountGroup.objects.filter(parent__isnull=False).count(), 0)
        self.assertEqual(AccountGroup.objects.filter(parent__isnull=True).count(), 4)

    def test_reset_page_delete_all_clears_accounts(self):
        from pathlib import Path
        from django.conf import settings

        template = (Path(settings.BASE_DIR) / "templates/admin_tools/reset_test_data.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("Accounts vouchers", template)
        self.assertIn("delete_accounts", template)
        self.client.force_login(self.admin)
        post = self.client.post(
            reverse("reset_test_data"),
            {"confirm": "1", "delete_all": "1"},
        )
        self.assertEqual(post.status_code, 302)
        self.assertEqual(Voucher.objects.count(), 0)
        self.assertEqual(InvoiceNumberSeq.objects.count(), 0)
        self.assertEqual(Account.objects.count(), 0)


class ClientLedgerBackfillTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="backfill-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        self.group = ClientGroup.objects.create(name="BACKFILL GROUP")

    def _make_client(self, name, *, approved=True, pan="AAAAA1111A"):
        return Client.objects.create(
            client_name=name,
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.group,
            approval_status=Client.APPROVED if approved else Client.PENDING,
            pan=pan,
        )

    def test_backfill_creates_only_approved_with_zero_opening(self):
        from accounts.services import SYSTEM_GROUP_DEBTORS, backfill_approved_client_ledgers

        approved = self._make_client("BACKFILL APPROVED", pan="AAAAA1111B")
        pending = self._make_client("BACKFILL PENDING", approved=False, pan="AAAAA1111C")
        Account.objects.filter(client=approved).delete()
        Account.objects.filter(client=pending).delete()

        preview = backfill_approved_client_ledgers(dry_run=True)
        self.assertGreaterEqual(preview["created"], 1)
        self.assertGreaterEqual(preview["skipped_unapproved"], 1)
        self.assertFalse(Account.objects.filter(client=approved).exists())

        result = backfill_approved_client_ledgers(dry_run=False)
        self.assertGreaterEqual(result["created"], 1)
        ledger = Account.objects.get(client=approved)
        self.assertEqual(ledger.opening_balance, Decimal("0.00"))
        self.assertEqual(ledger.group.system_key, SYSTEM_GROUP_DEBTORS)
        self.assertFalse(Account.objects.filter(client=pending).exists())

    def test_superuser_post_creates_ledgers(self):
        approved = self._make_client("BACKFILL POST CLIENT", pan="AAAAA1111D")
        Account.objects.filter(client=approved).delete()
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("accounts_sync_client_ledgers"),
            {"confirm": "1"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Account.objects.filter(client=approved).exists())


class ReportLedgerNavigationTests(SimpleTestCase):
    def test_templates_and_script_support_keyboard_drilldown(self):
        from pathlib import Path

        from django.conf import settings

        base = Path(settings.BASE_DIR)
        pnl = (base / "templates/accounts/report_pnl.html").read_text(encoding="utf-8")
        bs = (base / "templates/accounts/report_balance_sheet.html").read_text(encoding="utf-8")
        ledger = (base / "templates/accounts/report_ledger.html").read_text(encoding="utf-8")
        tb = (base / "templates/accounts/report_trial_balance.html").read_text(encoding="utf-8")
        script = (base / "static/js/report_ledger_nav.js").read_text(encoding="utf-8")
        tb_script = (base / "static/js/tb_row_nav.js").read_text(encoding="utf-8")
        self.assertIn("js-report-ledger-nav", pnl)
        self.assertIn("js-report-ledger-nav", bs)
        self.assertIn("report_ledger_nav.js", pnl)
        self.assertIn("report_ledger_nav.js", bs)
        self.assertIn("js-tb-nav", tb)
        self.assertIn("tb_row_nav.js", tb)
        self.assertIn("Escape", tb)
        self.assertIn('event.key === "Enter"', tb_script)
        self.assertIn("ArrowLeft", script)
        self.assertIn("ArrowRight", script)
        self.assertIn("ArrowDown", script)
        self.assertIn('params.set("next"', script)
        self.assertIn("Escape", ledger)
        self.assertIn("report_return_url", ledger)
        self.assertIn("js-ledger-entry", ledger)
        voucher = (base / "templates/accounts/voucher_detail.html").read_text(encoding="utf-8")
        self.assertIn("Escape", voucher)
        self.assertIn("back_to_ledger", voucher)


class LedgerAccountPickerTests(TestCase):
    def test_ledger_account_field_is_typeable_filtered_select(self):
        from pathlib import Path

        from django.conf import settings

        from accounts.forms import LedgerReportForm

        form = LedgerReportForm()
        self.assertIn("js-filtered-select", form.fields["account"].widget.attrs["class"])
        template = (Path(settings.BASE_DIR) / "templates/accounts/report_ledger.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("filtered_select.js", template)
        script = (Path(settings.BASE_DIR) / "static/js/filtered_select.js").read_text(encoding="utf-8")
        self.assertIn("ArrowDown", script)
        self.assertIn("ArrowUp", script)
        self.assertIn("choose(option)", script)


class PaymentRecoveryTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            email="rec-admin@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        grp = ClientGroup.objects.create(name="REC GROUP")
        self.client_row = Client.objects.create(
            client_name="REC CLIENT",
            client_type="Individual",
            branch="Trivandrum",
            client_group=grp,
            approval_status=Client.APPROVED,
            pan="RECDE7777F",
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        self.bank = Account.objects.create(name="Recovery bank", group=cash_group)
        current = AccountGroup.objects.create(name="Other current assets Rec", parent=assets, main_head="assets")
        self.gst_payment = Account.objects.create(
            name="GST Payment Rec",
            group=current,
            is_pure_agent=True,
        )
        expenses = AccountGroup.objects.get(system_key="expenses")
        self.rent = Account.objects.create(name="Office rent Rec", group=expenses)
        self.party = sync_client_account(self.client_row)
        tg = TaskGroup.objects.create(name="GST REC", sort_order=1)
        self.master = TaskMaster.objects.create(
            task_group=tg,
            name="GSTR recovery",
            is_recurring=False,
            requires_verifier=False,
            requires_document_checker=False,
        )
        self.task = Task.objects.create(
            client=self.client_row,
            task_master=self.master,
            title="GSTR-3B Jul Rec",
            due_date=date(2026, 7, 20),
            period_key="2026-07",
            status=Task.STATUS_ASSIGNED,
        )

    def test_pure_agent_payment_requires_client(self):
        from accounts.services import build_payment_lines

        lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.gst_payment, "amount": Decimal("5180.00")}],
        )
        with self.assertRaises(ValidationError):
            post_voucher(
                voucher_type="payment",
                voucher_date=date(2026, 7, 10),
                lines=lines,
                actor=self.admin,
            )

    def test_ordinary_payment_does_not_need_client(self):
        from accounts.services import build_payment_lines

        lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.rent, "amount": Decimal("100.00")}],
        )
        voucher = post_voucher(
            voucher_type="payment",
            voucher_date=date(2026, 7, 10),
            lines=lines,
            actor=self.admin,
        )
        self.assertIsNone(voucher.on_behalf_of_id)

    def test_tagged_payment_lists_until_invoice_knocks_off(self):
        from accounts.services import (
            build_invoice_lines,
            build_payment_lines,
            pending_recovery_rows,
        )

        pay_lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.gst_payment, "amount": Decimal("5180.00")}],
        )
        payment = post_voucher(
            voucher_type="payment",
            voucher_date=date(2026, 7, 10),
            lines=pay_lines,
            actor=self.admin,
            on_behalf_of=self.party,
            on_behalf_task=self.task,
        )
        self.assertEqual(payment.on_behalf_of_id, self.party.pk)
        pending = pending_recovery_rows(party=self.party)
        self.assertEqual(len(pending), 1)
        self.assertEqual(pending[0]["pending"], Decimal("5180.00"))
        self.assertEqual(pending[0]["task_id"], self.task.pk)
        income = _posting_ledgers()
        income.gst_rate = Decimal("18.00")
        income.save(update_fields=["gst_rate"])
        items = [
            {"account": income, "amount": Decimal("1000.00")},
            {
                "account": self.gst_payment,
                "amount": Decimal("5180.00"),
                "recovery_line": pending[0]["id"],
            },
        ]
        invoice_lines = build_invoice_lines(
            voucher_type="sales",
            party=self.party,
            items=items,
            with_gst=False,
        )
        invoice = post_voucher(
            voucher_type="sales",
            voucher_date=date(2026, 7, 12),
            lines=invoice_lines,
            actor=self.admin,
            recovery_items=items,
        )
        self.assertEqual(invoice.recovery_allocations.count(), 1)
        self.assertEqual(pending_recovery_rows(party=self.party), [])

    def test_wrong_task_client_is_rejected(self):
        from accounts.services import build_payment_lines

        other = Client.objects.create(
            client_name="OTHER REC",
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.client_row.client_group,
            approval_status=Client.APPROVED,
            pan="RECDE8888G",
        )
        other_task = Task.objects.create(
            client=other,
            task_master=self.master,
            title="Other GSTR Rec",
            due_date=date(2026, 7, 21),
            period_key="2026-07-other",
            status=Task.STATUS_ASSIGNED,
        )
        lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.gst_payment, "amount": Decimal("100.00")}],
        )
        with self.assertRaises(ValidationError):
            post_voucher(
                voucher_type="payment",
                voucher_date=date(2026, 7, 10),
                lines=lines,
                actor=self.admin,
                on_behalf_of=self.party,
                on_behalf_task=other_task,
            )

    def _tagged_payment(self, amount=Decimal("5180.00"), task=None):
        from accounts.services import build_payment_lines

        lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.gst_payment, "amount": amount}],
        )
        return post_voucher(
            voucher_type="payment",
            voucher_date=date(2026, 7, 10),
            lines=lines,
            actor=self.admin,
            on_behalf_of=self.party,
            on_behalf_task=task,
        )

    def test_unbilled_queue_fees_recovery_and_both(self):
        from accounts.services import (
            UNBILLED_FEES,
            UNBILLED_RECOVERY,
            build_unbilled_queue,
            pending_recovery_rows,
        )

        self._tagged_payment(task=self.task)
        queue = build_unbilled_queue(tasks=[])
        self.assertEqual(len(queue), 1)
        self.assertEqual(queue[0]["kind"], UNBILLED_RECOVERY)
        self.assertEqual(queue[0]["ledger_label"], "GST Payment Rec")
        self.assertEqual(queue[0]["key"], f"recovery:{pending_recovery_rows()[0]['id']}")

        self.task.status = Task.STATUS_COMPLETE
        self.task.is_billable = True
        self.task.fees_amount = Decimal("1000.00")
        self.task.save(update_fields=["status", "is_billable", "fees_amount", "updated_at"])
        queue = build_unbilled_queue(tasks=[self.task])
        kinds = {item["key"]: item["kind"] for item in queue}
        self.assertEqual(kinds[f"task:{self.task.pk}"], UNBILLED_FEES)
        recovery_item = next(item for item in queue if item["kind"] == UNBILLED_RECOVERY)
        self.assertEqual(recovery_item["ledger_label"], "GST Payment Rec")
        self.assertEqual(recovery_item["recovery"], Decimal("5180.00"))

        other = Task.objects.create(
            client=self.client_row,
            task_master=self.master,
            title="Fees only Rec",
            due_date=date(2026, 8, 20),
            period_key="2026-08",
            status=Task.STATUS_COMPLETE,
            is_billable=True,
            fees_amount=Decimal("250.00"),
        )
        queue = build_unbilled_queue(tasks=[self.task, other])
        kinds = {item["key"]: item["kind"] for item in queue}
        self.assertEqual(kinds[f"task:{self.task.pk}"], UNBILLED_FEES)
        self.assertEqual(kinds[f"task:{other.pk}"], UNBILLED_FEES)
        self.assertTrue(any(item["kind"] == UNBILLED_RECOVERY for item in queue))
        from accounts.services import filter_unbilled_queue

        found = filter_unbilled_queue(queue, "GST Payment")
        self.assertTrue(any(item["kind"] == UNBILLED_RECOVERY for item in found))

    def test_recovery_only_invoice_from_unbilled_keys(self):
        from accounts.services import (
            create_invoices_from_unbilled_keys,
            pending_recovery_rows,
        )

        self._tagged_payment(task=self.task)
        pending = pending_recovery_rows()
        created = create_invoices_from_unbilled_keys(
            keys=[f"recovery:{pending[0]['id']}"],
            tasks=[],
            actor=self.admin,
            voucher_date=date(2026, 7, 12),
            with_gst=False,
        )
        self.assertEqual(len(created), 1)
        self.assertTrue(created[0].number.startswith("NGST/"))
        self.assertEqual(created[0].recovery_allocations.count(), 1)
        self.assertEqual(pending_recovery_rows(), [])
        credits = {ln.account_id: ln.credit for ln in created[0].lines.all() if ln.credit > 0}
        self.assertEqual(credits[self.gst_payment.pk], Decimal("5180.00"))

    def test_same_client_rows_raise_one_invoice(self):
        from accounts.services import create_invoices_from_unbilled_keys, pending_recovery_rows

        self.task.status = Task.STATUS_COMPLETE
        self.task.is_billable = True
        self.task.fees_amount = Decimal("1000.00")
        self.task.save(update_fields=["status", "is_billable", "fees_amount", "updated_at"])
        other = Task.objects.create(
            client=self.client_row,
            task_master=self.master,
            title="GSTR-1 Jul Rec",
            due_date=date(2026, 7, 21),
            period_key="2026-07-gstr1",
            status=Task.STATUS_COMPLETE,
            is_billable=True,
            fees_amount=Decimal("250.00"),
        )
        self._tagged_payment(task=self.task)
        pending = pending_recovery_rows()[0]
        created = create_invoices_from_unbilled_keys(
            keys=[f"task:{self.task.pk}", f"task:{other.pk}", f"recovery:{pending['id']}"],
            tasks=[self.task, other],
            actor=self.admin,
            voucher_date=date(2026, 7, 12),
            with_gst=False,
        )
        self.assertEqual(len(created), 1)
        credits = {ln.account_id: ln.credit for ln in created[0].lines.all() if ln.credit > 0}
        self.assertEqual(credits[self.gst_payment.pk], Decimal("5180.00"))
        self.assertEqual(sum((ln.credit for ln in created[0].lines.all() if ln.account_id != self.gst_payment.pk and ln.credit > 0), Decimal("0.00")), Decimal("1250.00"))
        self.assertEqual(created[0].tasks.count(), 2)
        self.assertEqual(created[0].recovery_allocations.count(), 1)
        self.task.refresh_from_db()
        other.refresh_from_db()
        self.assertTrue(self.task.bill_raised)
        self.assertTrue(other.bill_raised)
        self.assertEqual(pending_recovery_rows(), [])

    def test_different_client_rows_raise_separate_invoices(self):
        from accounts.services import create_invoices_from_unbilled_keys

        self.task.status = Task.STATUS_COMPLETE
        self.task.is_billable = True
        self.task.fees_amount = Decimal("1000.00")
        self.task.save(update_fields=["status", "is_billable", "fees_amount", "updated_at"])
        other_client = Client.objects.create(
            client_name="OTHER BILL CLIENT",
            client_type="Individual",
            branch="Trivandrum",
            client_group=self.client_row.client_group,
            approval_status=Client.APPROVED,
            pan="RECDE9999H",
        )
        other_task = Task.objects.create(
            client=other_client,
            task_master=self.master,
            title="Other fees Rec",
            due_date=date(2026, 8, 20),
            period_key="2026-08",
            status=Task.STATUS_COMPLETE,
            is_billable=True,
            fees_amount=Decimal("400.00"),
        )
        created = create_invoices_from_unbilled_keys(
            keys=[f"task:{self.task.pk}", f"task:{other_task.pk}"],
            tasks=[self.task, other_task],
            actor=self.admin,
            voucher_date=date(2026, 7, 12),
            with_gst=False,
        )
        self.assertEqual(len(created), 2)
        parties = {next(ln.account_id for ln in voucher.lines.all() if ln.debit > 0) for voucher in created}
        self.assertEqual(len(parties), 2)

    def test_unbilled_page_posts_recovery_row(self):
        from accounts.services import pending_recovery_rows

        self._tagged_payment()
        pending = pending_recovery_rows()[0]
        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("accounts_unbilled_tasks"),
            {
                "item": f"recovery:{pending['id']}",
                "invoice_date": "2026-07-12",
                "invoice_type": "ngst",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(pending_recovery_rows(), [])

    def test_unbilled_template_lists_fees_and_recovery(self):
        from pathlib import Path

        from django.conf import settings

        template = (Path(settings.BASE_DIR) / "templates/accounts/unbilled_tasks.html").read_text(
            encoding="utf-8"
        )
        self.assertIn('name="item"', template)
        self.assertIn("Type", template)
        self.assertIn("Ledger", template)
        self.assertIn("Recovery", template)
        self.assertIn("Fees", template)
        sidebar = (Path(settings.BASE_DIR) / "templates/includes/app_sidebar_accounts_nav.html").read_text(
            encoding="utf-8"
        )
        self.assertNotIn("Pending recoveries", sidebar)


class PaymentVoucherPermissionTests(TestCase):
    def setUp(self):
        from django.core.exceptions import PermissionDenied
        from django.test import RequestFactory

        from accounts.permissions import voucher_action_flags
        from accounts.views import voucher_create, voucher_delete, voucher_edit, voucher_list, voucher_print

        self.PermissionDenied = PermissionDenied
        self.factory = RequestFactory()
        self.voucher_action_flags = voucher_action_flags
        self.voucher_create = voucher_create
        self.voucher_delete = voucher_delete
        self.voucher_edit = voucher_edit
        self.voucher_list = voucher_list
        self.voucher_print = voucher_print
        self.admin = User.objects.create_superuser(
            email="payperm-admin@example.com",
            password="pass12345",
        )
        seed_chart_of_accounts()
        assets = AccountGroup.objects.get(system_key="assets")
        cash_group = AccountGroup.objects.create(
            name="Cash and Cash Equivalents",
            parent=assets,
            main_head="assets",
        )
        self.bank = Account.objects.create(name="PayPerm bank", group=cash_group)
        self.rent = Account.objects.create(
            name="PayPerm rent",
            group=AccountGroup.objects.get(system_key="expenses"),
        )

    def _user_with(self, *codenames):
        from django.contrib.auth.models import Permission

        n = User.objects.filter(email__startswith="payperm-staff").count()
        user = User.objects.create_user(
            email=f"payperm-staff{n}@example.com",
            password="pass12345",
        )
        perms = Permission.objects.filter(content_type__app_label="accounts", codename__in=codenames)
        self.assertEqual(perms.count(), len(codenames), list(perms.values_list("codename", flat=True)))
        user.user_permissions.add(*perms)
        user = User.objects.get(pk=user.pk)
        self.client.force_login(user)
        return user

    def _payment_post(self, narration="PayPerm rent"):
        return {
            "voucher_type": "payment",
            "date": "2026-10-18",
            "narration": narration,
            "party": str(self.bank.pk),
            "form-TOTAL_FORMS": "1",
            "form-INITIAL_FORMS": "0",
            "form-MIN_NUM_FORMS": "0",
            "form-MAX_NUM_FORMS": "1000",
            "form-0-account": str(self.rent.pk),
            "form-0-amount": "100.00",
        }

    def _posted_payment(self):
        lines = build_payment_lines(
            mode=self.bank,
            items=[{"account": self.rent, "amount": Decimal("100.00")}],
        )
        return post_voucher(
            voucher_type="payment",
            voucher_date=date(2026, 10, 18),
            narration="Posted PayPerm",
            lines=lines,
            actor=self.admin,
        )

    def _request(self, user, method, path, data=None):
        if method == "post":
            request = self.factory.post(path, data or {})
        else:
            request = self.factory.get(path, data or {})
        request.user = user
        return request

    def test_post_voucher_cannot_add_payment(self):
        user = self._user_with("post_voucher", "view_voucher")
        request = self._request(
            user,
            "post",
            reverse("accounts_voucher_create"),
            self._payment_post(),
        )
        with self.assertRaises(self.PermissionDenied):
            self.voucher_create(request)
        self.assertFalse(Voucher.objects.filter(voucher_type="payment").exists())

    def test_add_payment_voucher_creates(self):
        self._user_with("add_payment_voucher")
        response = self.client.post(
            reverse("accounts_voucher_create") + "?type=payment",
            self._payment_post(),
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Voucher.objects.filter(number__startswith="PV/").exists())

    def test_payment_list_flags_follow_separate_perms(self):
        user = self._user_with("add_payment_voucher")
        flags = self.voucher_action_flags(user, "payment")
        self.assertTrue(flags["can_add"])
        self.assertFalse(flags["can_change"])
        self.assertFalse(flags["can_delete"])
        self.assertFalse(flags["can_print"])
        self.assertEqual(flags["voucher_tabs"], [("payment", "Payment")])

    def test_change_and_delete_are_separate(self):
        voucher = self._posted_payment()
        adder = self._user_with("add_payment_voucher")
        with self.assertRaises(self.PermissionDenied):
            self.voucher_edit(
                self._request(
                    adder,
                    "post",
                    reverse("accounts_voucher_edit", args=[voucher.pk]),
                    self._payment_post("Updated"),
                ),
                pk=voucher.pk,
            )
        with self.assertRaises(self.PermissionDenied):
            self.voucher_delete(
                self._request(adder, "post", reverse("accounts_voucher_delete", args=[voucher.pk])),
                pk=voucher.pk,
            )
        self.assertTrue(Voucher.objects.filter(pk=voucher.pk).exists())

        self._user_with("change_payment_voucher")
        edited = self.client.post(
            reverse("accounts_voucher_edit", args=[voucher.pk]),
            self._payment_post("Updated"),
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(edited.status_code, 302)
        voucher.refresh_from_db()
        self.assertEqual(voucher.narration, "Updated")

        self._user_with("delete_payment_voucher")
        deleted = self.client.post(
            reverse("accounts_voucher_delete", args=[voucher.pk]),
            HTTP_HOST="127.0.0.1",
        )
        self.assertEqual(deleted.status_code, 302)
        self.assertFalse(Voucher.objects.filter(pk=voucher.pk).exists())

    def test_print_payment_voucher_is_required(self):
        voucher = self._posted_payment()
        viewer = self._user_with("view_voucher")
        with self.assertRaises(self.PermissionDenied):
            self.voucher_list(
                self._request(
                    viewer,
                    "get",
                    reverse("accounts_voucher_list"),
                    {"type": "payment", "export": "pdf"},
                )
            )
        with self.assertRaises(self.PermissionDenied):
            self.voucher_print(
                self._request(viewer, "get", reverse("accounts_voucher_print", args=[voucher.pk])),
                pk=voucher.pk,
            )

        printer = self._user_with("print_payment_voucher")
        pdf_ok = self.voucher_list(
            self._request(
                printer,
                "get",
                reverse("accounts_voucher_list"),
                {"type": "payment", "export": "pdf"},
            )
        )
        self.assertEqual(pdf_ok.status_code, 200)
        self.assertTrue(pdf_ok.content.startswith(b"%PDF-"))
        from pathlib import Path

        from django.conf import settings

        from accounts.permissions import can_print_voucher

        self.assertTrue(can_print_voucher(printer, "payment"))
        template = (Path(settings.BASE_DIR) / "templates/accounts/voucher_print.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("Print / Save PDF", template)
        detail = (Path(settings.BASE_DIR) / "templates/accounts/voucher_detail.html").read_text(
            encoding="utf-8"
        )
        self.assertIn("accounts_voucher_print", detail)

